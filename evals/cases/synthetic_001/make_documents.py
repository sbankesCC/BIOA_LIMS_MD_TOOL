"""Writes the synthetic documents of eval case synthetic_001 into its `documents/` folder.

The committed files were made with this script; run it again to rebuild them
(`uv run python evals/cases/synthetic_001/make_documents.py`). Every value is made up.
"""

from pathlib import Path

import docx
from openpyxl import Workbook

DOCUMENTS_FOLDER = Path(__file__).resolve().parent / "documents"
METHOD_FILE = "synthetic_001_method.docx"
WORKSHEET_FILE = "synthetic_001_worksheet.xlsx"

# The method's reagent table (header row first)
METHOD_REAGENTS = [
    ["Item", "Vendor", "Cat. No.", "Pack", "Storage"],
    ["Synthetic Coating Antibody", "Synthetic Vendor One", "SYN-1001", "1 mL vial", "-20 °C"],
    [
        "Synthetic Wash Buffer",
        "Synthetic Vendor Two",
        "SYN-2001, SYN-2002",
        "500 mL bottle",
        "2–8 °C",
    ],
]

# The worksheet's reagent sheet (header row first)
WORKSHEET_REAGENTS = [
    ["Item", "Storage", "Lot"],
    ["Synthetic Coating Antibody", "2–8 °C", "SYN-LOT-01"],
    ["Synthetic Detection Reagent", "-20 °C", "SYN-LOT-02"],
]


def make_method(path: Path) -> None:
    """Write the synthetic method: a reagent table and an equipment section."""
    document = docx.Document()
    document.add_heading("Synthetic Method SYN-EVAL-001 (synthetic)", level=0)
    document.add_heading("6 Materials", level=1)
    document.add_heading("6.2 Reagents", level=2)
    document.add_paragraph("The following synthetic reagents are used.")
    # One row per reagent, cells as written in the table above
    table = document.add_table(rows=len(METHOD_REAGENTS), cols=len(METHOD_REAGENTS[0]))
    for row_index, row in enumerate(METHOD_REAGENTS):
        for column_index, text in enumerate(row):
            table.cell(row_index, column_index).text = text
    document.add_heading("6.3 Equipment", level=2)
    document.add_paragraph("Synthetic Plate Reader (synthetic instrument).")
    document.save(str(path))


def make_worksheet(path: Path) -> None:
    """Write the synthetic worksheet: one sheet listing reagents with storage and lot."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Reagents"
    for row in WORKSHEET_REAGENTS:
        sheet.append(row)
    workbook.save(path)


if __name__ == "__main__":
    DOCUMENTS_FOLDER.mkdir(exist_ok=True)
    make_method(DOCUMENTS_FOLDER / METHOD_FILE)
    make_worksheet(DOCUMENTS_FOLDER / WORKSHEET_FILE)
