"""Word reader: paragraphs grouped by heading path, and each table as `table_rows` chunks.

Body paragraphs and tables are read in document order. Headers, footers and text boxes are not read.
"""

import re
from pathlib import Path

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.contracts.chunk import Chunk, ChunkKind
from app.ingest.chunking import (
    MAX_TABLE_ROWS_PER_CHUNK,
    MAX_TEXT_CHUNK_CHARACTERS,
    ChunkIdCounter,
    format_row,
    group_paragraphs,
    is_blank_row,
    split_table_rows,
)

# Word's built-in heading styles ("Heading 1", "Heading 2", ...); "Title" counts as level 0
HEADING_STYLE_PATTERN = re.compile(r"^Heading (\d+)$")
TITLE_STYLE = "Title"

# How the heading path is shown in a chunk's `section`
SECTION_SEPARATOR = " > "


def read_docx_chunks(path: Path, document_id: str) -> list[Chunk]:
    """Read a Word file into paragraph and table chunks, in reading order."""
    builder = WordChunkBuilder(document_id)
    for block in docx.Document(str(path)).iter_inner_content():
        if isinstance(block, Paragraph):
            builder.add_paragraph(block)
        else:
            builder.add_table(block)
    builder.flush_paragraphs()
    return builder.chunks


def heading_level(paragraph: Paragraph) -> int | None:
    """Return the heading level of a paragraph, or None for body text."""
    style = paragraph.style
    style_name = style.name if style is not None else None
    if style_name == TITLE_STYLE:
        return 0
    match = HEADING_STYLE_PATTERN.match(style_name or "")
    return int(match.group(1)) if match else None


def table_rows(table: Table) -> list[str]:
    """Return a table's non-blank rows as lines of text (line breaks in a cell become spaces)."""
    rows: list[str] = []
    for row in table.rows:
        cells = [" ".join(cell.text.split("\n")) for cell in row.cells]
        if not is_blank_row(cells):
            rows.append(format_row(cells))
    return rows


class WordChunkBuilder:
    """Collects chunks while walking a Word body: tracks the heading path and pending text."""

    def __init__(self, document_id: str) -> None:
        """Start with no heading, no pending paragraphs and no chunks."""
        self.document_id = document_id
        self.chunk_ids = ChunkIdCounter(document_id)
        self.headings: list[tuple[int, str]] = []
        self.pending: list[str] = []
        self.chunks: list[Chunk] = []

    def section(self) -> str | None:
        """Return the current heading path, e.g. `6 Materials > 6.2 Reagents`, or None."""
        if not self.headings:
            return None
        return SECTION_SEPARATOR.join(text for _, text in self.headings)

    def add_paragraph(self, paragraph: Paragraph) -> None:
        """Add body text to the current section, or start a new section at a heading."""
        text = paragraph.text
        level = heading_level(paragraph)
        if level is None:
            if text.strip():
                self.pending.append(text)
            return
        # A heading closes the current section and replaces headings at its level or deeper
        self.flush_paragraphs()
        if text.strip():
            self.headings = [item for item in self.headings if item[0] < level]
            self.headings.append((level, text.strip()))

    def add_table(self, table: Table) -> None:
        """Turn a table into one `table_rows` chunk, or header-repeated groups if very large."""
        # Text before the table comes first, to keep reading order
        self.flush_paragraphs()
        rows = table_rows(table)
        if not rows:
            return
        for group in split_table_rows(rows, MAX_TABLE_ROWS_PER_CHUNK):
            self.add_chunk("table_rows", "\n".join(group))

    def flush_paragraphs(self) -> None:
        """Turn the pending paragraphs of the current section into paragraph chunks."""
        for group in group_paragraphs(self.pending, MAX_TEXT_CHUNK_CHARACTERS):
            self.add_chunk("paragraph", "\n".join(group))
        self.pending = []

    def add_chunk(self, kind: ChunkKind, text: str) -> None:
        """Append one text chunk in the current section."""
        chunk = Chunk(
            chunk_id=self.chunk_ids.next_id(),
            document_id=self.document_id,
            page=None,
            section=self.section(),
            sheet=None,
            row_range=None,
            kind=kind,
            is_scanned=False,
            text=text,
            image_ref=None,
        )
        self.chunks.append(chunk)
