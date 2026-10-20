"""DFD loader (ARCHITECTURE §4.10): a copy of the template with one row per included record.

Only `output.value` is written; excluded records are left out. Number fields are written as
numbers; a value that is not a plain number stays text (validation flagged it as WRONG_TYPE).
Fields the tool does not fill (`blank_for_now`, `system`) are exported blank.
"""

import shutil
from collections.abc import Sequence
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.worksheet.worksheet import Worksheet

from app.config.dfd_template import SHEET_NAME, DfdTemplate
from app.config.field_rules import FieldRule, FieldRules
from app.contracts.record import StoredRecord
from app.validation.field_checks import is_right_type

HEADER_ROW = 1
FIRST_DATA_ROW = 2
NUMBER_TYPES = ("decimal", "integer")

# openpyxl reads text starting with "=" as a formula; such values are forced to plain text
FORMULA_START = "="
TEXT_DATA_TYPE = "s"

# A loader cell value: text, a number, or blank
LoaderValue = str | int | float | None


def write_loader(
    template: DfdTemplate, field_rules: FieldRules, records: Sequence[StoredRecord], target: Path
) -> int:
    """Copy the template to `target` and fill one row per included record; return the row count."""
    # The template itself is only read: the rows go into a copy
    shutil.copyfile(template.path, target)
    workbook = load_workbook(target)
    sheet = workbook[SHEET_NAME]
    # Each template column is filled from the field rule with the same column name
    rules_by_column = {rule.column: rule for rule in field_rules.fields}
    headers = read_header_columns(sheet)
    included = [record for record in records if record.exclusion is None]
    for row_number, record in enumerate(included, start=FIRST_DATA_ROW):
        for column_number, header in headers.items():
            cell = find_cell(sheet, row_number, column_number)
            write_cell(cell, find_loader_value(rules_by_column[header], record))
    workbook.save(target)
    return len(included)


def read_header_columns(sheet: Worksheet) -> dict[int, str]:
    """Return the template's header text by column number."""
    return {
        cell.column: str(cell.value)
        for cell in sheet[HEADER_ROW]
        if isinstance(cell, Cell) and cell.value is not None
    }


def find_export_text(rule: FieldRule, record: StoredRecord) -> str | None:
    """Return the text a field exports: its output value, or None for blank and unfilled fields."""
    stored = record.fields.get(rule.key)
    if rule.role != "extracted" or stored is None or not stored.output.value:
        return None
    return stored.output.value


def find_loader_value(rule: FieldRule, record: StoredRecord) -> LoaderValue:
    """Return a field's loader cell value: a number for a number field holding a plain number."""
    text = find_export_text(rule, record)
    if text is not None and rule.type in NUMBER_TYPES and is_right_type(rule, text):
        return float(text) if "." in text else int(text)
    return text


def find_cell(sheet: Worksheet, row: int, column: int) -> Cell:
    """Return a writable cell; a merged cell cannot hold a value, so it is refused."""
    cell = sheet.cell(row=row, column=column)
    if isinstance(cell, MergedCell):
        raise ValueError(
            f"Cell {cell.coordinate} of sheet '{sheet.title}' is merged and cannot hold a value. "
            "Unmerge the cells in the template."
        )
    return cell


def write_cell(cell: Cell, value: LoaderValue) -> None:
    """Write a value to a cell, keeping text that starts with "=" as text, not a formula."""
    cell.value = value
    if isinstance(value, str) and value.startswith(FORMULA_START):
        cell.data_type = TEXT_DATA_TYPE
