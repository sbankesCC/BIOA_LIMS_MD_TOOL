"""Lookup lists workbook (`lookup_lists.xlsx`): one sheet of allowed values per lookup field.

Sheets starting with "_" (`_Instructions`, hidden `_metadata`) are reserved, not lookup lists.
"""

from collections.abc import Sequence
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel

from app.config.errors import ConfigError, require_file

FILE_NAME = "lookup_lists.xlsx"
HEADER = "Lookup Value"
METADATA_SHEET = "_metadata"
METADATA_KEYS = ("version_id", "based_on_version", "exported_at", "exported_by")
RESERVED_PREFIX = "_"


class LookupMetadata(BaseModel):
    """The hidden `_metadata` sheet: which version this workbook is and where it came from."""

    version_id: str
    based_on_version: str | None = None
    exported_at: str | None = None
    exported_by: str | None = None


class LookupLists(BaseModel):
    """All lookup lists: values per sheet name, in workbook order, plus metadata."""

    metadata: LookupMetadata
    sheets: dict[str, list[str]]

    @property
    def version(self) -> str:
        """Return the lookup lists version id from `_metadata`."""
        return self.metadata.version_id


def load_lookup_lists(path: Path) -> LookupLists:
    """Load the lookup workbook and check headers, values and metadata."""
    require_file(path)
    workbook = load_workbook(path, data_only=True)
    worksheets = workbook.worksheets
    problems: list[str] = []
    sheets: dict[str, list[str]] = {}
    # Every non-reserved sheet is a lookup list
    for worksheet in worksheets:
        if worksheet.title.startswith(RESERVED_PREFIX):
            continue
        values, sheet_problems = read_lookup_sheet(worksheet)
        sheets[worksheet.title] = values
        problems.extend(sheet_problems)
    metadata = read_metadata(worksheets, problems)
    if problems or metadata is None:
        raise ConfigError(problems)
    return LookupLists(metadata=metadata, sheets=sheets)


def read_lookup_sheet(worksheet: Worksheet) -> tuple[list[str], list[str]]:
    """Read one lookup sheet: header in A1, then one text value per row with no gaps."""
    name = f"{FILE_NAME}: sheet '{worksheet.title}'"
    rows = [row[0] if row else None for row in worksheet.iter_rows(max_col=1, values_only=True)]
    if not rows or rows[0] != HEADER:
        return [], [f"{name}: cell A1 must be '{HEADER}'."]
    values: list[str] = []
    problems: list[str] = []
    # Trailing empty rows are ignored; an empty row before a value is a gap
    last_filled = max((index for index, cell in enumerate(rows) if cell is not None), default=0)
    for row_number, cell in enumerate(rows[1 : last_filled + 1], start=2):
        if not isinstance(cell, str) or not cell.strip():
            problems.append(f"{name}: row {row_number} must be a text value, not blank.")
            continue
        values.append(cell)
    return values, problems


def read_metadata(worksheets: Sequence[Worksheet], problems: list[str]) -> LookupMetadata | None:
    """Read the `_metadata` sheet (key in column A, value in B); add problems if it is wrong."""
    metadata_sheets = [sheet for sheet in worksheets if sheet.title == METADATA_SHEET]
    if not metadata_sheets:
        problems.append(f"{FILE_NAME}: the hidden '{METADATA_SHEET}' sheet is missing.")
        return None
    entries: dict[str, str | None] = {}
    for row in metadata_sheets[0].iter_rows(max_col=2, values_only=True):
        key = row[0] if row else None
        value = row[1] if len(row) > 1 else None
        if isinstance(key, str) and key in METADATA_KEYS:
            entries[key] = None if value is None else str(value)
    # The version id is what runs record and what stale uploads are checked against
    version_id = entries.get("version_id")
    if not version_id:
        problems.append(f"{FILE_NAME}: '{METADATA_SHEET}' has no version_id.")
        return None
    return LookupMetadata(
        version_id=version_id,
        based_on_version=entries.get("based_on_version"),
        exported_at=entries.get("exported_at"),
        exported_by=entries.get("exported_by"),
    )
