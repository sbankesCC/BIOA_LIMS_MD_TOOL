"""Lookup list upload checks (ARCHITECTURE §4.11) and the diff against the current version.

Problems block the upload. A value longer than its field's length limit is only a warning
(user decision 2026-10-20): records with it still get TOO_LONG in validation.
"""

from collections.abc import Sequence
from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook, load_workbook
from pydantic import BaseModel

from app.admin.reports import UploadFinding, make_problem, make_warning
from app.config.field_rules import FieldRules
from app.config.lookup_lists import (
    HEADER,
    RESERVED_PREFIX,
    LookupLists,
    LookupMetadata,
    read_metadata,
)
from app.validation.field_checks import is_control_character

# Finding codes (shown to the admin with the message; the API passes them through)
FILE_UNREADABLE = "FILE_UNREADABLE"
METADATA_MISSING = "METADATA_MISSING"
STALE_VERSION = "STALE_VERSION"
USED_SHEET_MISSING = "USED_SHEET_MISSING"
HEADER_WRONG = "HEADER_WRONG"
BLANK_ROW = "BLANK_ROW"
NOT_TEXT = "NOT_TEXT"
CONTROL_CHARACTERS = "CONTROL_CHARACTERS"
EXTRA_SPACES = "EXTRA_SPACES"
DUPLICATE_VALUE = "DUPLICATE_VALUE"
CASE_ONLY_DUPLICATE = "CASE_ONLY_DUPLICATE"
TOO_LONG = "TOO_LONG"

# Which problem codes each passed check stands for (a check passes when none of them is found)
PASSED_CHECKS: dict[str, tuple[str, ...]] = {
    "Based on the current version ({version})": (FILE_UNREADABLE, METADATA_MISSING, STALE_VERSION),
    "All expected sheets and headers present": (FILE_UNREADABLE, USED_SHEET_MISSING, HEADER_WRONG),
    "No blank rows, extra spaces or control characters": (
        FILE_UNREADABLE,
        BLANK_ROW,
        NOT_TEXT,
        CONTROL_CHARACTERS,
        EXTRA_SPACES,
    ),
    "No duplicate values": (FILE_UNREADABLE, DUPLICATE_VALUE, CASE_ONLY_DUPLICATE),
}


@dataclass(frozen=True)
class UploadedSheet:
    """One lookup sheet as uploaded: whether A1 is the header, and column A cells by row number."""

    name: str
    has_header: bool
    cells: Sequence[tuple[int, object]]


class SheetChange(BaseModel):
    """Values added and removed on one sheet, compared with the current version."""

    sheet: str
    added: list[str]
    removed: list[str]
    is_new_sheet: bool = False
    is_removed_sheet: bool = False


class LookupUploadReport(BaseModel):
    """The check report of one upload: passed checks, problems, warnings and changes."""

    file_name: str
    current_version: str
    based_on_version: str | None
    passed_checks: list[str]
    problems: list[UploadFinding]
    warnings: list[UploadFinding]
    changes: list[SheetChange]

    def has_problems(self) -> bool:
        """Return True if any problem blocks applying this upload."""
        return bool(self.problems)


class LookupCheckResult(BaseModel):
    """The report, plus the uploaded lists when no problem was found."""

    report: LookupUploadReport
    lookup_lists: LookupLists | None


def check_lookup_upload(
    file_name: str, content: bytes, field_rules: FieldRules, current: LookupLists
) -> LookupCheckResult:
    """Check an uploaded lookup workbook against the field rules and the current version."""
    workbook = open_workbook(content)
    if workbook is None:
        problem = make_problem(
            FILE_UNREADABLE,
            "The file could not be read as an Excel workbook (.xlsx). Upload "
            "the workbook downloaded from this page.",
        )
        return LookupCheckResult(
            report=build_report(file_name, current, None, [problem], []), lookup_lists=None
        )
    findings: list[UploadFinding] = []
    metadata = check_metadata(workbook, current.version, findings)
    sheets = read_uploaded_sheets(workbook)
    findings += check_used_sheets(sheets, field_rules)
    limits = find_sheet_limits(field_rules)
    for sheet in sheets:
        findings += check_sheet(sheet, limits.get(sheet.name))
    uploaded = {sheet.name: read_sheet_values(sheet) for sheet in sheets}
    report = build_report(
        file_name, current, metadata, findings, diff_sheets(current.sheets, uploaded)
    )
    # The lists are kept only when nothing blocks the upload
    lookup_lists = None
    if metadata is not None and not report.has_problems():
        lookup_lists = LookupLists(metadata=metadata, sheets=uploaded)
    return LookupCheckResult(report=report, lookup_lists=lookup_lists)


def open_workbook(content: bytes) -> Workbook | None:
    """Open uploaded bytes as a workbook; return None if they are not a readable .xlsx."""
    try:
        return load_workbook(BytesIO(content), data_only=True)
    except Exception:
        # Any reader error (not a zip, not a workbook, broken parts) means "unreadable"
        return None


def build_report(
    file_name: str,
    current: LookupLists,
    metadata: LookupMetadata | None,
    findings: Sequence[UploadFinding],
    changes: list[SheetChange],
) -> LookupUploadReport:
    """Sort the findings into problems and warnings and list the checks that passed."""
    found_codes = {finding.code for finding in findings if finding.severity == "problem"}
    passed = [
        label.format(version=current.version)
        for label, codes in PASSED_CHECKS.items()
        if not found_codes.intersection(codes)
    ]
    return LookupUploadReport(
        file_name=file_name,
        current_version=current.version,
        based_on_version=metadata.version_id if metadata is not None else None,
        passed_checks=passed,
        problems=[finding for finding in findings if finding.severity == "problem"],
        warnings=[finding for finding in findings if finding.severity == "warning"],
        changes=changes,
    )


def check_metadata(
    workbook: Workbook, current_version: str, findings: list[UploadFinding]
) -> LookupMetadata | None:
    """Read `_metadata`; add a problem if it is missing or based on an older version."""
    metadata = read_metadata(workbook.worksheets, [])
    if metadata is None:
        findings.append(
            make_problem(
                METADATA_MISSING,
                "The hidden _metadata sheet is missing or has no version_id. Download the current "
                "workbook, make your edits in it, and upload it again.",
            )
        )
        return None
    # A workbook downloaded before another version was applied would undo that version's edits
    if metadata.version_id != current_version:
        findings.append(
            make_problem(
                STALE_VERSION,
                f"This file is based on version {metadata.version_id}, but the current version "
                f"is {current_version}. Download the current workbook, make your edits again, "
                "and upload it.",
            )
        )
    return metadata


def read_uploaded_sheets(workbook: Workbook) -> list[UploadedSheet]:
    """Read every lookup sheet (names not starting with "_"), column A only."""
    sheets: list[UploadedSheet] = []
    for worksheet in workbook.worksheets:
        if worksheet.title.startswith(RESERVED_PREFIX):
            continue
        cells = [
            row[0] if row else None for row in worksheet.iter_rows(max_col=1, values_only=True)
        ]
        has_header = bool(cells) and cells[0] == HEADER
        # Empty rows after the last value are ignored; empty rows before it are blank rows
        last_filled = max(
            (index for index, cell in enumerate(cells) if not is_blank(cell)), default=0
        )
        numbered = list(enumerate(cells[1 : last_filled + 1], start=2))
        sheets.append(UploadedSheet(name=worksheet.title, has_header=has_header, cells=numbered))
    return sheets


def is_blank(cell: object) -> bool:
    """Return True for an empty cell or one holding only spaces."""
    return cell is None or (isinstance(cell, str) and not cell.strip())


def check_used_sheets(
    sheets: Sequence[UploadedSheet], field_rules: FieldRules
) -> list[UploadFinding]:
    """Return a problem for each sheet a field uses that is missing (renamed or deleted)."""
    names = {sheet.name for sheet in sheets}
    return [
        make_problem(
            USED_SHEET_MISSING,
            f"The sheet '{rule.lookup_sheet}' is missing. The field {rule.column} uses it, so it "
            "cannot be renamed or deleted. Put it back with its exact name and upload again.",
            sheet=rule.lookup_sheet,
        )
        for rule in field_rules.fields
        if rule.lookup_sheet is not None and rule.lookup_sheet not in names
    ]


def find_sheet_limits(field_rules: FieldRules) -> dict[str, int]:
    """Return the length limit of each lookup sheet (the smallest of the fields using it)."""
    limits: dict[str, int] = {}
    for rule in field_rules.fields:
        if rule.lookup_sheet is None or rule.max_length is None:
            continue
        limits[rule.lookup_sheet] = min(
            rule.max_length, limits.get(rule.lookup_sheet, rule.max_length)
        )
    return limits


def check_sheet(sheet: UploadedSheet, limit: int | None) -> list[UploadFinding]:
    """Check one sheet: header, then every value, then duplicates."""
    if not sheet.has_header:
        return [
            make_problem(
                HEADER_WRONG,
                f"Cell A1 must be '{HEADER}'. Put the header back and upload again.",
                sheet=sheet.name,
                row=1,
            )
        ]
    findings: list[UploadFinding] = []
    for row, cell in sheet.cells:
        findings += check_cell(sheet.name, row, cell, limit)
    findings += find_duplicates(sheet)
    return findings


def check_cell(sheet: str, row: int, cell: object, limit: int | None) -> list[UploadFinding]:
    """Check one value: not blank, text, no control characters or edge spaces, within the limit."""
    if is_blank(cell):
        return [make_problem(BLANK_ROW, "This row is blank. Delete the empty row.", sheet, row)]
    if not isinstance(cell, str):
        return [
            make_problem(
                NOT_TEXT,
                f"{cell!r} is not saved as text. Format the cell as text and type it again.",
                sheet,
                row,
            )
        ]
    findings: list[UploadFinding] = []
    if any(is_control_character(character) for character in cell):
        message = f"“{cell}” holds a control character (such as a line break). Remove it."
        findings.append(make_problem(CONTROL_CHARACTERS, message, sheet, row))
    elif cell != cell.strip():
        message = f"“{cell}” starts or ends with a space. Remove the space."
        findings.append(make_problem(EXTRA_SPACES, message, sheet, row))
    # Over the limit is reported but allowed: records with the value are flagged TOO_LONG
    if limit is not None and len(cell) > limit:
        message = (
            f"“{cell}” has {len(cell)} characters; the field's limit is {limit}. It can be "
            "applied, but records with this value will be flagged as too long."
        )
        findings.append(make_warning(TOO_LONG, message, sheet, row))
    return findings


def find_duplicates(sheet: UploadedSheet) -> list[UploadFinding]:
    """Return a problem for each value that repeats an earlier one, exactly or by case only."""
    first_by_key: dict[str, str] = {}
    findings: list[UploadFinding] = []
    for row, cell in sheet.cells:
        if not isinstance(cell, str) or is_blank(cell):
            continue
        # The first value with a key is kept; later ones are compared with it
        earlier = first_by_key.get(cell.casefold())
        if earlier is None:
            first_by_key[cell.casefold()] = cell
            continue
        if earlier == cell:
            message = f"“{cell}” is listed twice. Delete one of them."
            findings.append(make_problem(DUPLICATE_VALUE, message, sheet.name, row))
        else:
            message = (
                f"“{cell}” duplicates “{earlier}” (differs only by capitalization). "
                "Keep one of them."
            )
            findings.append(make_problem(CASE_ONLY_DUPLICATE, message, sheet.name, row))
    return findings


def read_sheet_values(sheet: UploadedSheet) -> list[str]:
    """Return a sheet's text values in order (blank and non-text cells are reported, not kept)."""
    return [cell for _, cell in sheet.cells if isinstance(cell, str) and not is_blank(cell)]


def diff_sheets(current: dict[str, list[str]], uploaded: dict[str, list[str]]) -> list[SheetChange]:
    """Return the sheets whose values differ from the current version, in workbook order."""
    changes: list[SheetChange] = []
    names = list(uploaded) + [name for name in current if name not in uploaded]
    for name in names:
        before = current.get(name, [])
        after = uploaded.get(name, [])
        added = [value for value in after if value not in before]
        removed = [value for value in before if value not in after]
        is_new_sheet = name not in current
        is_removed_sheet = name not in uploaded
        if added or removed or is_new_sheet or is_removed_sheet:
            changes.append(
                SheetChange(
                    sheet=name,
                    added=added,
                    removed=removed,
                    is_new_sheet=is_new_sheet,
                    is_removed_sheet=is_removed_sheet,
                )
            )
    return changes
