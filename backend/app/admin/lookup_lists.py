"""Lookup list admin (ARCHITECTURE §4.11): download, upload with checks and diff, apply, restore.

The seed workbook in the client config is never written. Admin uploads are stored under
`STORAGE_DIR/lookup_lists/`, with their version rows in the review database. New runs use the
current version, or the seed while no admin version has been applied.
"""

from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, JsonValue

from app.admin.errors import AdminError
from app.admin.lookup_checks import LookupUploadReport, check_lookup_upload
from app.admin.versions import (
    apply_version,
    find_version,
    make_admin_event,
    make_file_ref,
    next_upload_id,
    resolve_file_ref,
)
from app.config.bundle import ConfigBundle
from app.config.lookup_lists import (
    FILE_NAME,
    METADATA_KEYS,
    METADATA_SHEET,
    LookupLists,
    LookupMetadata,
    load_lookup_lists,
)
from app.logging_setup import log_event
from app.storage.review_db import DataVersion, DataVersionKind, ReviewDatabase

LOOKUP_FOLDER = "lookup_lists"
KIND: DataVersionKind = "lookup"

# Admin events (admin_events table)
UPLOADED_EVENT = "lookup_lists_uploaded"
APPLIED_EVENT = "lookup_version_applied"
RESTORED_EVENT = "lookup_version_restored"


class LookupUploadResult(BaseModel):
    """An upload's check report, and the pending version created when it has no problems."""

    report: LookupUploadReport
    version: DataVersion | None


class LookupListAdmin:
    """Lookup list versions: the seed from the client config plus admin uploads."""

    def __init__(self, database: ReviewDatabase, storage_dir: Path, bundle: ConfigBundle) -> None:
        """Use the review database, STORAGE_DIR, and the seed and field rules from the bundle."""
        self.database = database
        self.storage_dir = storage_dir
        self.field_rules = bundle.field_rules
        self.seed = bundle.lookup_lists
        self.seed_path = bundle.config_dir / FILE_NAME

    # ----- Reading versions -----

    def find_current_lookup_lists(self) -> LookupLists:
        """Return the lookup lists a new run uses: the current admin version, or the seed."""
        current = self.database.read_current_data_version(KIND)
        return self.seed if current is None else self.load_version(current)

    def load_lookup_lists(self, version_id: str) -> LookupLists:
        """Return the lookup lists of one version (e.g. the version a run captured at its start)."""
        if version_id == self.seed.version:
            return self.seed
        return self.load_version(find_version(self.database, KIND, version_id))

    def load_version(self, version: DataVersion) -> LookupLists:
        """Load a stored upload; its version name comes from the database, not the file."""
        lookup_lists = load_lookup_lists(resolve_file_ref(self.storage_dir, version.file_ref))
        metadata = LookupMetadata(
            version_id=version.version_id or version.upload_id,
            based_on_version=version.based_on_version,
        )
        return lookup_lists.model_copy(update={"metadata": metadata})

    def list_versions(self) -> list[DataVersion]:
        """Return every admin upload in upload order (the seed is not listed)."""
        return self.database.list_data_versions(KIND)

    # ----- Download -----

    def download_workbook(self, user: str, at: datetime, version_id: str | None = None) -> bytes:
        """Return a version's workbook (the current one by default) with `_metadata` filled in."""
        name = version_id or self.find_current_lookup_lists().version
        if name == self.seed.version:
            source, based_on = self.seed_path, self.seed.metadata.based_on_version
        else:
            version = find_version(self.database, KIND, name)
            source = resolve_file_ref(self.storage_dir, version.file_ref)
            name = version.version_id or version.upload_id
            based_on = version.based_on_version
        # The file is edited in memory only: the seed and stored uploads are never written
        workbook = load_workbook(source)
        metadata = LookupMetadata(
            version_id=name, based_on_version=based_on, exported_at=at.isoformat(), exported_by=user
        )
        write_metadata(workbook, metadata)
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    # ----- Upload -----

    def upload_workbook(
        self, file_name: str, content: bytes, user: str, at: datetime
    ) -> LookupUploadResult:
        """Check an uploaded workbook; without problems it is stored as a pending version."""
        current = self.find_current_lookup_lists()
        result = check_lookup_upload(file_name, content, self.field_rules, current)
        version = None
        if result.lookup_lists is not None:
            version = self.store_upload(content, result.report, user, at)
        event = make_admin_event(
            UPLOADED_EVENT,
            user,
            at,
            upload_id=version.upload_id if version is not None else None,
            file_name=file_name,
            problems=len(result.report.problems),
            warnings=len(result.report.warnings),
        )
        self.database.add_admin_event(event)
        return LookupUploadResult(report=result.report, version=version)

    def store_upload(
        self, content: bytes, report: LookupUploadReport, user: str, at: datetime
    ) -> DataVersion:
        """Save the uploaded file under STORAGE_DIR and add its pending version row."""
        with self.database.transaction():
            upload_id = next_upload_id(self.list_versions())
            path = self.storage_dir / LOOKUP_FOLDER / f"{upload_id}.xlsx"
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(content)
            version = DataVersion(
                upload_id=upload_id,
                file_ref=make_file_ref(self.storage_dir, path),
                based_on_version=report.based_on_version,
                uploaded_by=user,
                uploaded_at=at,
                summary=summarize_report(report),
            )
            self.database.add_data_version(KIND, version)
        return version

    # ----- Apply and restore -----

    def apply_version(self, name: str, user: str, at: datetime) -> DataVersion:
        """Make a pending upload current as the next `vN`, or restore an older version."""
        with self.database.transaction():
            version = find_version(self.database, KIND, name)
            if version.status == "pending":
                self.check_not_stale(version)
            applied = apply_version(self.database, KIND, version, user, at)
            event_name = RESTORED_EVENT if applied.is_restore else APPLIED_EVENT
            event = make_admin_event(
                event_name,
                user,
                at,
                upload_id=applied.version.upload_id,
                version_id=applied.version.version_id,
            )
            self.database.add_admin_event(event)
        log_event("lookup_version_applied")
        return applied.version

    def check_not_stale(self, version: DataVersion) -> None:
        """Refuse a pending upload whose base is no longer the current version."""
        current_version = self.find_current_lookup_lists().version
        if version.based_on_version != current_version:
            raise AdminError(
                "UPLOAD_STALE",
                f"This upload is based on version {version.based_on_version}, but the current "
                f"version is now {current_version}. Download the current workbook, make your "
                "edits again, and upload it.",
            )


def write_metadata(workbook: Workbook, metadata: LookupMetadata) -> None:
    """Replace the hidden `_metadata` sheet's rows (key in column A, value in B)."""
    if METADATA_SHEET in workbook.sheetnames:
        sheet = workbook[METADATA_SHEET]
        sheet.delete_rows(1, sheet.max_row)
    else:
        sheet = workbook.create_sheet(METADATA_SHEET)
    sheet.sheet_state = "hidden"
    for key in METADATA_KEYS:
        sheet.append([key, getattr(metadata, key)])


def summarize_report(report: LookupUploadReport) -> dict[str, JsonValue]:
    """Return what the version history shows: file name, value counts and the full report."""
    return {
        "file_name": report.file_name,
        "added_values": sum(len(change.added) for change in report.changes),
        "removed_values": sum(len(change.removed) for change in report.changes),
        "report": report.model_dump(mode="json"),
    }
