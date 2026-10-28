"""API end-to-end test with the fake LLM, every step through the HTTP API.

Upload, processing, review, changes, re-check, Verify, export and download.
"""

from io import BytesIO
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from app.config.bundle import ConfigBundle
from tests.fixtures.fake_llm import TABLE_CHUNK_ID
from tests.fixtures.make_documents import REAGENT_TABLE, WORD_FILE
from tests.integration.api_helpers import (
    USER,
    XLSX_TYPE,
    ApiHarness,
    happy_replies,
    make_settings,
    open_api,
    read_file,
)

RUN_ID = "run_0001"


def exclude(change_id: str, record_id: str) -> dict[str, Any]:
    """Return an exclude_record change."""
    return {
        "type": "exclude_record",
        "change_id": change_id,
        "record_id": record_id,
        "reason": "synthetic: not needed",
    }


def find_shared_issue(view: dict[str, Any]) -> dict[str, Any]:
    """Return an open field issue that both records have on the same field with the same value."""
    first, second = view["rows"]
    for cell in first["cells"]:
        other = next(c for c in second["cells"] if c["field"] == cell["field"])
        for issue in cell["issues"]:
            other_codes = [i["code"] for i in other["issues"] if i["status"] == "open"]
            if (
                issue["status"] == "open"
                and not issue["individual_review"]
                and issue["code"] in other_codes
                and cell["value"] == other["value"]
            ):
                return issue
    raise AssertionError("The synthetic run has no shared issue to preview.")


def check_processed_run(api: ApiHarness) -> None:
    """Check the run reached review with both records, its document and the fake models."""
    run = api.client.get(f"/runs/{RUN_ID}").json()
    assert run["status"] == "in_review"
    assert run["current_step"] is None
    assert run["counts"]["records"] == 2
    assert run["models"]["extraction"] == "fake-model"
    assert [document["file_name"] for document in run["documents"]] == [WORD_FILE]
    runs = api.client.get("/runs").json()
    assert [summary["run_id"] for summary in runs] == [RUN_ID]
    assert runs[0]["counts"]["records"] == 2


def check_review_view(api: ApiHarness, bundle: ConfigBundle) -> dict[str, Any]:
    """Check one row per record; placeholder fields show the placeholder, value empty."""
    view = api.client.get(f"/runs/{RUN_ID}/records").json()
    assert [row["record_id"] for row in view["rows"]] == ["rec_0001", "rec_0002"]
    assert view["rows"][0]["name"] == REAGENT_TABLE[1][0]
    placeholders = [cell for cell in view["rows"][0]["cells"] if cell["is_placeholder"]]
    assert placeholders
    for cell in placeholders:
        assert cell["value"] is None
        assert cell["display_value"] == bundle.field_rules.placeholder_value
    return view


def check_record_and_source(api: ApiHarness) -> None:
    """Check a record has its sections, issues and sources, and its source chunk the text."""
    record = api.client.get(f"/runs/{RUN_ID}/records/rec_0001").json()
    assert record["record"]["record_id"] == "rec_0001"
    assert record["issues"]
    assert [source["chunk_id"] for source in record["sources"]] == [TABLE_CHUNK_ID]
    chunk = api.client.get(f"/runs/{RUN_ID}/chunks/{TABLE_CHUNK_ID}").json()
    assert REAGENT_TABLE[1][0] in chunk["text"]
    assert chunk["location"]["file_name"] == WORD_FILE


def check_bulk_preview(api: ApiHarness, view: dict[str, Any]) -> None:
    """Check a field issue both records share previews both records."""
    issue = find_shared_issue(view)
    response = api.client.get(
        f"/runs/{RUN_ID}/bulk-preview",
        params={"mode": "change_all", "issue_id": issue["issue_id"]},
    )
    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview["code"] == issue["code"]
    assert preview["count"] == 2
    assert [match["record_id"] for match in preview["matches"]] == ["rec_0001", "rec_0002"]


def check_draft_export(api: ApiHarness, bundle: ConfigBundle) -> None:
    """Check an unverified run exports a draft whose loader holds the one included record."""
    response = api.client.post(f"/runs/{RUN_ID}/exports")
    assert response.status_code == 201, response.text
    export = response.json()
    assert export["is_draft"] is True
    loader = api.client.get(f"/runs/{RUN_ID}/exports/{export['export_id']}/loader")
    assert loader.status_code == 200
    assert loader.headers["content-type"] == XLSX_TYPE
    assert "draft" in loader.headers["content-disposition"]
    sheet = load_workbook(BytesIO(loader.content)).active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))
    assert list(rows[0]) == bundle.dfd_template.headers
    assert len(rows) == 2
    audit = api.client.get(f"/runs/{RUN_ID}/exports/{export['export_id']}/audit")
    assert audit.status_code == 200
    assert len(load_workbook(BytesIO(audit.content)).sheetnames) == 5


@pytest.mark.usefixtures("blank_environment")
def test_upload_to_export_download(
    tmp_path: Path, bundle: ConfigBundle, synthetic_documents: dict[str, Path]
) -> None:
    """A run goes from upload to a verified export through the API alone."""
    with open_api(make_settings(tmp_path), happy_replies(bundle)) as api:
        started = api.start_run([read_file(synthetic_documents, WORD_FILE)])
        assert started.status_code == 202, started.text
        assert started.json() == {"run_id": RUN_ID, "retry_of": None}
        check_processed_run(api)
        view = check_review_view(api, bundle)
        check_record_and_source(api)
        check_bulk_preview(api, view)
        # Exclude one record; a resent change ID is ignored
        saved = api.post_changes(RUN_ID, [exclude("chg_0001", "rec_0002")]).json()
        assert saved["saved_change_ids"] == ["chg_0001"]
        resent = api.post_changes(RUN_ID, [exclude("chg_0001", "rec_0002")]).json()
        assert resent["ignored_change_ids"] == ["chg_0001"]
        assert api.client.get(f"/runs/{RUN_ID}").json()["has_pending_changes"] is True
        recheck = api.client.post(f"/runs/{RUN_ID}/recheck").json()
        assert recheck["applied_change_ids"] == ["chg_0001"]
        assert recheck["counts"]["excluded"] == 1
        check_draft_export(api, bundle)
        # The remaining record has red issues, so Verify is blocked
        blocked = api.client.post(f"/runs/{RUN_ID}/verify").json()
        assert blocked["is_verified"] is False
        assert blocked["blocking_errors"]
        # With both records excluded, Verify accepts the run warnings and sets the label
        api.post_changes(RUN_ID, [exclude("chg_0002", "rec_0001")])
        verified = api.client.post(f"/runs/{RUN_ID}/verify").json()
        assert verified["is_verified"] is True
        export = api.client.post(f"/runs/{RUN_ID}/exports").json()
        assert export["is_draft"] is False
        assert export["number"] == 2
        run = api.client.get(f"/runs/{RUN_ID}").json()
        assert run["is_verified"] is True
        assert run["verified_by"] == USER
        assert len(api.client.get(f"/runs/{RUN_ID}/exports").json()) == 2
