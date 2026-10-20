"""Export service tests: loader round-trip, audit sheets, file names, re-check, export record.

Synthetic records only; field keys, columns and types are read from the client config.
"""

import dataclasses
import hashlib
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.config.bundle import ConfigBundle
from app.config.dfd_template import SHEET_NAME
from app.config.field_rules import FieldRule
from app.contracts.chunk import Chunk
from app.contracts.export import ExportInfo
from app.contracts.extraction import Alternative, Evidence, ExtractedField
from app.contracts.record import FieldOutput, StoredField, StoredRecord
from app.export import audit_workbook
from app.export.audit_rows import find_written_text
from app.export.service import ExportContext, export_run, make_export_file_names
from app.review.errors import RunNotInReviewError
from app.review.service import ReviewContext
from app.storage.review_db import ReviewDatabase
from app.storage.run_store import RunFolder, RunStore
from app.validation.duplicates import find_name_rule
from tests.unit.review_helpers import (
    NOW,
    RUN_ID,
    USER,
    exclude_record,
    make_context,
    make_relaxed_config,
    recheck,
    save,
    set_value,
    start_review,
)
from tests.unit.validation_helpers import find_rule, make_record, set_output, update_field

CHUNK_ID = "doc01-c001"
QUOTE = "synthetic quote from the method"
DOCUMENT_NAME = "synthetic method.docx"
RUN_NAME = "synthetic run"


@pytest.fixture
def context(bundle: ConfigBundle) -> ReviewContext:
    """Return a review context on the relaxed config."""
    return make_context(make_relaxed_config(bundle))


@pytest.fixture
def run_folder(tmp_path: Path) -> RunFolder:
    """Create the run's folder with one synthetic document and one chunk."""
    folder = RunStore(tmp_path / "storage").create_run(RUN_ID)
    folder.save_upload(DOCUMENT_NAME, b"synthetic document")
    chunk = Chunk(
        chunk_id=CHUNK_ID,
        document_id="doc01",
        page=None,
        section="Synthetic section",
        sheet=None,
        row_range=None,
        kind="paragraph",
        is_scanned=False,
        text=QUOTE,
        image_ref=None,
    )
    folder.write_chunks([chunk])
    return folder


def number_rule(bundle: ConfigBundle, field_type: str) -> FieldRule:
    """Return the first extracted field of a number type."""
    return find_rule(bundle.field_rules, lambda rule: rule.type == field_type)


def unfilled_rules(bundle: ConfigBundle) -> list[FieldRule]:
    """Return the fields the tool does not fill (blank_for_now and system)."""
    return [rule for rule in bundle.field_rules.fields if rule.role != "extracted"]


def make_records(bundle: ConfigBundle) -> list[StoredRecord]:
    """Return three synthetic records; the first has a decimal value and a cited description."""
    first = set_output(
        make_record(bundle, "rec_0001", "Syn A"), number_rule(bundle, "decimal").key, "2.5"
    )
    # A value that sneaks into an unfilled field must still be exported blank
    first = set_output(first, unfilled_rules(bundle)[0].key, "synthetic", origin="blank_for_now")
    cited = text_rule(bundle)
    extraction = ExtractedField(
        status="found", value="synthetic", evidence=[Evidence(chunk_id=CHUNK_ID, quote=QUOTE)]
    )
    first = update_field(
        first, cited.key, first.fields[cited.key].model_copy(update={"extraction": extraction})
    )
    return [
        first,
        make_record(bundle, "rec_0002", "Syn B"),
        make_record(bundle, "rec_0003", "Syn C"),
    ]


def text_rule(bundle: ConfigBundle) -> FieldRule:
    """Return a free-text extracted field that is not the record name."""
    name_key = find_name_rule(bundle.field_rules).key
    return find_rule(
        bundle.field_rules,
        lambda rule: rule.type == "text" and rule.lookup_sheet is None and rule.key != name_key,
    )


def start_export_run(
    tmp_path: Path, bundle: ConfigBundle, context: ReviewContext
) -> ReviewDatabase:
    """Start review on the synthetic records and exclude the second one."""
    database = start_review(tmp_path, bundle, context, make_records(bundle))
    save(database, context, exclude_record("chg_0001", "rec_0002", "synthetic: not a consumable"))
    recheck(database, context)
    return database


def export(
    database: ReviewDatabase, bundle: ConfigBundle, context: ReviewContext, run_folder: RunFolder
) -> ExportInfo:
    """Export the run for the test user."""
    export_context = ExportContext(
        review=context, template=bundle.dfd_template, models={"extraction": "synthetic-model"}
    )
    return export_run(database, export_context, run_folder, RUN_ID, USER, NOW)


def open_sheet(path: Path, title: str) -> Worksheet:
    """Open one sheet of a written workbook."""
    sheet = load_workbook(path)[title]
    assert isinstance(sheet, Worksheet)
    return sheet


def read_rows(sheet: Worksheet) -> list[tuple[object, ...]]:
    """Return every row of a sheet as values."""
    return list(sheet.iter_rows(values_only=True))


# ----- Loader -----


def test_loader_round_trips_with_headers_and_rows(
    tmp_path: Path, bundle: ConfigBundle, context: ReviewContext, run_folder: RunFolder
) -> None:
    """The loader has the template headers and one row per included record, as exported."""
    database = start_export_run(tmp_path, bundle, context)
    info = export(database, bundle, context, run_folder)
    rows = read_rows(open_sheet(run_folder.resolve_ref(info.loader_file), SHEET_NAME))
    headers = bundle.dfd_template.headers
    assert list(rows[0]) == headers
    name_index = headers.index(find_name_rule(bundle.field_rules).column)
    # The excluded record is left out
    assert [row[name_index] for row in rows[1:]] == ["Syn A", "Syn C"]
    records = {r.record_id: r for r in database.read_records(RUN_ID)}
    for rule in bundle.field_rules.fields:
        value = rows[2][headers.index(rule.column)]
        expected = records["rec_0003"].fields[rule.key].output.value
        if rule.role != "extracted":
            assert value is None
        elif rule.type in ("decimal", "integer"):
            assert value == int(expected or "") and isinstance(value, int)
        else:
            assert value == expected


def test_loader_writes_numbers_as_numbers_and_unfilled_fields_blank(
    tmp_path: Path, bundle: ConfigBundle, context: ReviewContext, run_folder: RunFolder
) -> None:
    """A decimal is a float cell; blank_for_now and system fields are empty."""
    database = start_export_run(tmp_path, bundle, context)
    info = export(database, bundle, context, run_folder)
    rows = read_rows(open_sheet(run_folder.resolve_ref(info.loader_file), SHEET_NAME))
    headers = bundle.dfd_template.headers
    decimal_value = rows[1][headers.index(number_rule(bundle, "decimal").column)]
    assert decimal_value == 2.5 and isinstance(decimal_value, float)
    integer_value = rows[1][headers.index(number_rule(bundle, "integer").column)]
    assert integer_value == 1 and isinstance(integer_value, int)
    for rule in unfilled_rules(bundle):
        assert rows[1][headers.index(rule.column)] is None


def test_unreadable_number_stays_text(
    tmp_path: Path, bundle: ConfigBundle, context: ReviewContext, run_folder: RunFolder
) -> None:
    """A number field holding text that is not a plain number is exported as written."""
    rule = number_rule(bundle, "decimal")
    database = start_export_run(tmp_path, bundle, context)
    save(database, context, set_value("chg_0002", rule.key, "two", reason="synthetic edit"))
    info = export(database, bundle, context, run_folder)
    rows = read_rows(open_sheet(run_folder.resolve_ref(info.loader_file), SHEET_NAME))
    assert rows[1][bundle.dfd_template.headers.index(rule.column)] == "two"


def test_text_starting_with_equals_is_not_a_formula(
    tmp_path: Path, bundle: ConfigBundle, context: ReviewContext, run_folder: RunFolder
) -> None:
    """A value starting with "=" is written as text, never as a formula."""
    rule = text_rule(bundle)
    database = start_export_run(tmp_path, bundle, context)
    save(database, context, set_value("chg_0002", rule.key, "=1+1", reason="synthetic edit"))
    info = export(database, bundle, context, run_folder)
    sheet = open_sheet(run_folder.resolve_ref(info.loader_file), SHEET_NAME)
    cell = sheet.cell(row=2, column=bundle.dfd_template.headers.index(rule.column) + 1)
    assert (cell.value, cell.data_type) == ("=1+1", "s")


def test_template_is_never_written(
    tmp_path: Path, bundle: ConfigBundle, context: ReviewContext, run_folder: RunFolder
) -> None:
    """Exporting fills a copy; the template file is unchanged."""
    before = hashlib.sha256(bundle.dfd_template.path.read_bytes()).hexdigest()
    export(start_export_run(tmp_path, bundle, context), bundle, context, run_folder)
    assert hashlib.sha256(bundle.dfd_template.path.read_bytes()).hexdigest() == before


# ----- Re-check, file names, export record -----


def test_export_applies_pending_changes_first(
    tmp_path: Path, bundle: ConfigBundle, context: ReviewContext, run_folder: RunFolder
) -> None:
    """A pending change is applied (re-check) before the loader is written."""
    rule = text_rule(bundle)
    database = start_export_run(tmp_path, bundle, context)
    save(database, context, set_value("chg_0002", rule.key, "synthetic edited", reason="fix"))
    info = export(database, bundle, context, run_folder)
    rows = read_rows(open_sheet(run_folder.resolve_ref(info.loader_file), SHEET_NAME))
    assert rows[1][bundle.dfd_template.headers.index(rule.column)] == "synthetic edited"
    assert database.read_changes(RUN_ID, "pending") == []


def test_file_names_are_numbered_with_draft_until_verified(
    tmp_path: Path, bundle: ConfigBundle, context: ReviewContext, run_folder: RunFolder
) -> None:
    """Export 1 of an unverified run is a draft; export 2 after verifying is not."""
    database = start_export_run(tmp_path, bundle, context)
    first = export(database, bundle, context, run_folder)
    database.set_verified(RUN_ID, USER, NOW)
    second = export(database, bundle, context, run_folder)
    assert Path(first.loader_file).name == f"{RUN_NAME} – export 1 – draft.xlsx"
    assert Path(first.audit_file).name == f"{RUN_NAME} – export 1 – audit – draft.xlsx"
    assert Path(second.loader_file).name == f"{RUN_NAME} – export 2.xlsx"
    assert Path(second.audit_file).name == f"{RUN_NAME} – export 2 – audit.xlsx"
    assert (first.is_draft, second.is_draft) == (True, False)
    assert run_folder.resolve_ref(second.audit_file).is_file()


def test_unsafe_characters_in_run_name_are_replaced() -> None:
    """Characters Windows does not allow in file names become "_"."""
    names = make_export_file_names('synthetic: a/b "run"', 3, is_draft=False)
    assert names.loader == "synthetic_ a_b _run_ – export 3.xlsx"


def test_export_is_recorded(
    tmp_path: Path, bundle: ConfigBundle, context: ReviewContext, run_folder: RunFolder
) -> None:
    """The export is stored with its number and files, with an audit event."""
    database = start_export_run(tmp_path, bundle, context)
    info = export(database, bundle, context, run_folder)
    assert database.list_exports(RUN_ID) == [info]
    assert (info.export_id, info.number, info.exported_by) == ("exp_0001", 1, USER)
    event = database.read_audit_events(RUN_ID)[-1]
    assert (event.event, event.details["is_draft"]) == ("export_created", True)


def test_export_needs_a_run_in_review(
    tmp_path: Path, bundle: ConfigBundle, context: ReviewContext, run_folder: RunFolder
) -> None:
    """A run that is not in review cannot be exported."""
    database = start_export_run(tmp_path, bundle, context)
    database.set_run_status(RUN_ID, "failed")
    with pytest.raises(RunNotInReviewError):
        export(database, bundle, context, run_folder)


# ----- Audit workbook -----


def test_audit_workbook_has_five_sheets_with_columns(
    tmp_path: Path, bundle: ConfigBundle, context: ReviewContext, run_folder: RunFolder
) -> None:
    """The audit file has the CONTRACTS §8 sheets in order, each with its header row."""
    database = start_export_run(tmp_path, bundle, context)
    path = run_folder.resolve_ref(export(database, bundle, context, run_folder).audit_file)
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Summary", "Values", "Edits", "Excluded", "Run issues"]
    name_column = find_name_rule(bundle.field_rules).column
    expected_headers = {
        "Summary": audit_workbook.SUMMARY_HEADERS,
        "Values": audit_workbook.make_value_headers(name_column),
        "Edits": audit_workbook.make_edit_headers(name_column),
        "Excluded": audit_workbook.EXCLUDED_HEADERS,
        "Run issues": audit_workbook.RUN_ISSUE_HEADERS,
    }
    for title, headers in expected_headers.items():
        assert list(read_rows(open_sheet(path, title))[0]) == headers
    assert len(audit_workbook.make_value_headers(name_column)) == 14
    assert len(audit_workbook.make_edit_headers(name_column)) == 11


def test_audit_values_cover_every_field_of_included_records(
    tmp_path: Path, bundle: ConfigBundle, context: ReviewContext, run_folder: RunFolder
) -> None:
    """Values has one row per included record and field, with sources resolved."""
    database = start_export_run(tmp_path, bundle, context)
    path = run_folder.resolve_ref(export(database, bundle, context, run_folder).audit_file)
    rows = read_rows(open_sheet(path, "Values"))[1:]
    field_count = len(bundle.field_rules.fields)
    assert len(rows) == 2 * field_count
    assert {row[0] for row in rows} == {"rec_0001", "rec_0003"}
    cited = next(row for row in rows if row[0] == "rec_0001" and row[2] == text_rule(bundle).column)
    # Value as written sits just before the exported value
    assert (cited[3], cited[4]) == ("synthetic", "synthetic")
    assert (cited[6], cited[7], cited[8]) == (
        DOCUMENT_NAME,
        f"{DOCUMENT_NAME}, Synthetic section",
        QUOTE,
    )


def multi_value_rule(bundle: ConfigBundle) -> FieldRule:
    """Return the first extracted field that holds several values."""
    return find_rule(bundle.field_rules, lambda rule: rule.multi_value)


def stored_extraction(extraction: ExtractedField) -> StoredField:
    """Return a stored field holding only this extraction (output left empty)."""
    return StoredField(extraction=extraction, output=FieldOutput(origin="extracted"))


def test_written_value_joins_a_multi_value_with_the_field_separator(bundle: ConfigBundle) -> None:
    """A list as written is joined with the field's separator."""
    rule = multi_value_rule(bundle)
    stored = stored_extraction(ExtractedField(status="found", value=["syn-1", "syn-2"]))
    assert find_written_text(rule, stored) == f"syn-1{rule.separator}syn-2"


def test_written_value_of_a_conflict_lists_each_alternative_on_its_own_line(
    bundle: ConfigBundle,
) -> None:
    """A multi-value conflict shows one alternative per line, each list joined as usual."""
    rule = multi_value_rule(bundle)
    evidence = [Evidence(chunk_id=CHUNK_ID, quote=QUOTE)]
    alternatives = [
        Alternative(value=["syn-1", "syn-2"], evidence=evidence),
        Alternative(value=["syn-3"], evidence=evidence),
    ]
    stored = stored_extraction(ExtractedField(status="conflicting", alternatives=alternatives))
    assert find_written_text(rule, stored) == f"syn-1{rule.separator}syn-2\nsyn-3"


def test_written_value_of_a_single_value_conflict_is_one_per_line(bundle: ConfigBundle) -> None:
    """A conflict on a single-value field lists its alternatives one per line too."""
    rule = find_rule(bundle.field_rules, lambda rule: not rule.multi_value)
    evidence = [Evidence(chunk_id=CHUNK_ID, quote=QUOTE)]
    alternatives = [
        Alternative(value="syn-a", evidence=evidence),
        Alternative(value="syn-b", evidence=evidence),
    ]
    stored = stored_extraction(ExtractedField(status="conflicting", alternatives=alternatives))
    assert find_written_text(rule, stored) == "syn-a\nsyn-b"


def test_written_value_is_blank_when_not_found_or_not_extracted(bundle: ConfigBundle) -> None:
    """A not-found field and a field the tool does not fill have no value as written."""
    rule = text_rule(bundle)
    not_found = stored_extraction(ExtractedField(status="not_found"))
    assert find_written_text(rule, not_found) is None
    assert find_written_text(unfilled_rules(bundle)[0], None) is None


def test_audit_lists_edits_exclusions_summary_and_run_issues(
    tmp_path: Path, bundle: ConfigBundle, context: ReviewContext, run_folder: RunFolder
) -> None:
    """Edits, Excluded, Summary and Run issues hold the run's history and settings."""
    no_reference = dataclasses.replace(context, reference=None)
    database = start_export_run(tmp_path, bundle, no_reference)
    path = run_folder.resolve_ref(export(database, bundle, no_reference, run_folder).audit_file)
    edits = read_rows(open_sheet(path, "Edits"))[1:]
    assert [(row[0], row[1], row[2], row[4]) for row in edits] == [
        ("chg_0001", "rec_0002", "Syn B", "exclude_record")
    ]
    excluded = read_rows(open_sheet(path, "Excluded"))[1:]
    assert [(row[0], row[1], row[3], row[5]) for row in excluded] == [
        ("rec_0002", "Syn B", USER, "synthetic: not a consumable")
    ]
    summary = {str(row[0]): row[1] for row in read_rows(open_sheet(path, "Summary"))[1:]}
    assert (summary["Run ID"], summary["Export type"]) == (RUN_ID, "Draft")
    assert summary["Lookup lists version"] == bundle.lookup_lists.version
    assert summary["Model extraction"] == "synthetic-model"
    assert "Document doc01" in summary
    run_issue_codes = [row[0] for row in read_rows(open_sheet(path, "Run issues"))[1:]]
    assert "REFERENCE_DATA_MISSING" in run_issue_codes
