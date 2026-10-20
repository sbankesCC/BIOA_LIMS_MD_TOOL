"""Config consistency: field rules and seed lookup lists must match BioA's source workbook.

The source workbook (Data info + List Values sheets) is the only source of truth. Every allowed
difference is a named rule with its reason below, and each one must actually be used: an
unused rule means the config or the workbook changed and the list needs a review.
"""

import re
from dataclasses import dataclass, field
from typing import Any

import pytest
from openpyxl import load_workbook

from app.config.bundle import ConfigBundle
from app.config.field_rules import FieldRule, fields_in_group
from tests.conftest import SOURCE_WORKBOOK

# Allowed differences between the config and the source workbook, with their reasons
ALLOWED_DIFFERENCES = {
    "fields_to_add_are_blank_for_now": (
        "Rows under 'Fields to add:' (not system generated) are not filled by the tool yet: "
        "role blank_for_now, exported blank (decided with BioA)."
    ),
    "mandatory_in_condition_group_is_conditional": (
        "Data info marks the amount, amount unit and container count fields Mandatory, but its "
        "note says amount + unit OR containers is mandatory: read as conditional."
    ),
    "blank_lookup_entry_removed": (
        "A blank entry in a List Values column is not a value; it was left out of the seed."
    ),
    "lookup_value_whitespace_trimmed": (
        "A leading line break in a List Values cell is a copy artefact; it was removed."
    ),
}

# Source workbook layout
DATA_INFO_SHEET = "Data info"
LIST_VALUES_SHEET = "List Values"
FIELDS_TO_ADD_MARKER = "Fields to add:"
DEFAULT_NOTE_PATTERN = re.compile(r"Default\s*=\s*(.+)$")
DATA_TYPES = {"Char": "text", "Decimal": "decimal", "Number": "integer"}


@dataclass
class SourceField:
    """One field row from the Data info sheet, with the 'Fields to add' position."""

    column: str
    field_type: str
    data_type: str
    mandatory_optional: str
    field_input: str
    character_limit: int | None
    notes: str | None
    is_field_to_add: bool


@dataclass
class Comparison:
    """Differences found, and which allowed-difference rules were used to explain others."""

    differences: list[str] = field(default_factory=list)
    used_rules: set[str] = field(default_factory=set)


def read_source_fields() -> list[SourceField]:
    """Read the Data info sheet into field rows, in order."""
    worksheet = load_workbook(SOURCE_WORKBOOK, data_only=True)[DATA_INFO_SHEET]
    rows = list(worksheet.iter_rows(min_row=2, values_only=True))
    fields: list[SourceField] = []
    is_field_to_add = False
    for row in rows:
        if row[0] is None:
            continue
        # The marker row splits current fields from fields BioA wants to add
        if str(row[0]).strip() == FIELDS_TO_ADD_MARKER:
            is_field_to_add = True
            continue
        fields.append(source_field_from_row(row, is_field_to_add))
    return fields


def source_field_from_row(row: tuple[Any, ...], is_field_to_add: bool) -> SourceField:
    """Build a SourceField from one Data info row."""
    return SourceField(
        column=str(row[0]),
        field_type=str(row[1]),
        data_type=str(row[2]),
        mandatory_optional=str(row[4]),
        field_input=str(row[5]),
        character_limit=row[6] if isinstance(row[6], int) else None,
        notes=None if row[7] is None else str(row[7]),
        is_field_to_add=is_field_to_add,
    )


def read_source_lists() -> dict[str, list[str]]:
    """Read the List Values sheet: one column per lookup list, header in row 1."""
    worksheet = load_workbook(SOURCE_WORKBOOK, data_only=True)[LIST_VALUES_SHEET]
    columns = list(worksheet.iter_cols(values_only=True))
    lists: dict[str, list[str]] = {}
    for column in columns:
        if column[0] is None:
            continue
        lists[str(column[0])] = [str(cell) for cell in column[1:] if cell is not None]
    return lists


def expected_role(source: SourceField, comparison: Comparison) -> str:
    """Return the role the source implies, noting when an allowed difference applies."""
    if source.field_type == "System Generated":
        return "system"
    if source.is_field_to_add:
        comparison.used_rules.add("fields_to_add_are_blank_for_now")
        return "blank_for_now"
    return "extracted"


def expected_requirement(source: SourceField, grouped: set[str], rule: FieldRule) -> str | None:
    """Return the requirement the source implies; None when the group rule explains it."""
    if source.mandatory_optional == "Mandatory":
        return None if rule.key in grouped else "required"
    return "optional"


def expected_type(source: SourceField) -> str:
    """Return the field type: Boolean input wins, otherwise the Data Type column."""
    if source.field_input == "Boolean":
        return "boolean"
    return DATA_TYPES[source.data_type]


def expected_default(source: SourceField) -> str | None:
    """Return the default from a 'Default = X' note, if there is one."""
    match = DEFAULT_NOTE_PATTERN.search(source.notes or "")
    return match.group(1).strip() if match else None


def compare_field(
    source: SourceField, rule: FieldRule, grouped: set[str], comparison: Comparison
) -> None:
    """Compare one source row with its field rule, recording each difference."""
    checks: dict[str, tuple[object, object]] = {
        "role": (expected_role(source, comparison), rule.role),
        "type": (expected_type(source), rule.type),
        "max_length": (source.character_limit, rule.max_length),
        "has lookup sheet": (source.field_input == "List", rule.lookup_sheet == rule.column),
        "default": (expected_default(source), rule.default),
    }
    requirement = expected_requirement(source, grouped, rule)
    if requirement is None:
        comparison.used_rules.add("mandatory_in_condition_group_is_conditional")
        requirement = "conditional"
    checks["requirement"] = (requirement, rule.requirement)
    for name, (expected, actual) in checks.items():
        if expected != actual:
            comparison.differences.append(
                f"{rule.column}: {name} is {actual!r}, source gives {expected!r}"
            )


def clean_source_values(source_values: list[str], comparison: Comparison) -> list[str]:
    """Clean the source values by the allowed rules and return them for comparison."""
    cleaned: list[str] = []
    for value in source_values:
        if not value.strip():
            comparison.used_rules.add("blank_lookup_entry_removed")
            continue
        if value != value.strip():
            comparison.used_rules.add("lookup_value_whitespace_trimmed")
        cleaned.append(value.strip())
    return cleaned


@pytest.fixture(scope="module")
def field_comparison(bundle: ConfigBundle) -> Comparison:
    """Compare every field rule with its Data info row."""
    comparison = Comparison()
    grouped = {
        key
        for group in bundle.field_rules.condition_groups.values()
        for key in fields_in_group(group)
    }
    for source, rule in zip(read_source_fields(), bundle.field_rules.fields, strict=False):
        compare_field(source, rule, grouped, comparison)
    return comparison


@pytest.fixture(scope="module")
def lookup_comparison(bundle: ConfigBundle) -> Comparison:
    """Compare every seed lookup sheet with its List Values column."""
    comparison = Comparison()
    source_lists = read_source_lists()
    for sheet_name, seed_values in bundle.lookup_lists.sheets.items():
        source_values = source_lists.get(sheet_name, [])
        cleaned = clean_source_values(source_values, comparison)
        if cleaned != seed_values:
            comparison.differences.append(f"Lookup sheet '{sheet_name}' differs from List Values")
    return comparison


def test_field_list_and_order_match_source(bundle: ConfigBundle) -> None:
    """The field rules have the Data info fields, in the same order."""
    source_columns = [source.column for source in read_source_fields()]
    assert bundle.field_rules.columns() == source_columns


def test_field_settings_match_source(field_comparison: Comparison) -> None:
    """Type, requirement, limits, list/boolean input, role and defaults match Data info."""
    assert field_comparison.differences == []


def test_every_source_list_is_a_lookup_sheet(bundle: ConfigBundle) -> None:
    """Each List Values column has a lookup sheet of the same name."""
    assert set(read_source_lists()) <= set(bundle.lookup_lists.sheets)


def test_seed_lookup_values_match_source(lookup_comparison: Comparison) -> None:
    """The seed lookup values equal List Values, after the allowed clean-ups."""
    assert lookup_comparison.differences == []


def test_every_allowed_difference_is_used(
    field_comparison: Comparison, lookup_comparison: Comparison
) -> None:
    """Each allowed difference explains something; an unused one needs a review."""
    used = field_comparison.used_rules | lookup_comparison.used_rules
    assert used == set(ALLOWED_DIFFERENCES)
