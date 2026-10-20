"""Lookup list admin tests: download, upload checks, warnings, diff, apply, restore, seed safety.

Uploads start from the downloaded workbook and get synthetic edits; sheet names, limits and
seed values are read from the client config, never written here.
"""

import hashlib
from collections.abc import Callable
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.admin import lookup_checks
from app.admin.errors import AdminError
from app.admin.lookup_checks import diff_sheets
from app.admin.lookup_lists import LookupListAdmin, LookupUploadResult
from app.config.bundle import ConfigBundle
from app.config.field_rules import FieldRule
from app.config.lookup_lists import HEADER, METADATA_SHEET
from app.storage.review_db import ReviewDatabase
from tests.unit.review_helpers import NOW

ADMIN = "synthetic.admin"
SYNTHETIC_VALUE = "Synthetic value"

# An edit applied to the downloaded workbook before it is uploaded
WorkbookEdit = Callable[[Workbook], None]


@pytest.fixture
def admin(tmp_path: Path, bundle: ConfigBundle) -> LookupListAdmin:
    """Return a lookup admin with a new review database and storage folder."""
    database = ReviewDatabase(tmp_path / "review.sqlite")
    return LookupListAdmin(database, tmp_path / "storage", bundle)


def roomy_rule(bundle: ConfigBundle) -> FieldRule:
    """Return a lookup field with at least two values and room for the synthetic values."""
    return next(
        rule
        for rule in bundle.field_rules.fields
        if rule.lookup_sheet is not None
        and len(bundle.lookup_lists.sheets[rule.lookup_sheet]) >= 2
        and (rule.max_length or 0) >= len(SYNTHETIC_VALUE) + 5
    )


def roomy_sheet(workbook: Workbook, bundle: ConfigBundle) -> Worksheet:
    """Return the worksheet of `roomy_rule` in a workbook."""
    sheet_name = roomy_rule(bundle).lookup_sheet
    assert sheet_name is not None
    sheet = workbook[sheet_name]
    assert isinstance(sheet, Worksheet)
    return sheet


def next_empty_row(sheet: Worksheet) -> int:
    """Return the row after the last value in column A."""
    filled = [cell.row for cell in sheet["A"] if cell.value is not None]
    return max(filled) + 1


def append_values(sheet: Worksheet, *values: str | int) -> None:
    """Write values under the last value of column A."""
    for value in values:
        sheet.cell(row=next_empty_row(sheet), column=1, value=value)


def edit_workbook(content: bytes, edit: WorkbookEdit) -> bytes:
    """Apply an edit to workbook bytes and return the new bytes."""
    workbook = load_workbook(BytesIO(content))
    edit(workbook)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def upload_edited(admin: LookupListAdmin, edit: WorkbookEdit) -> LookupUploadResult:
    """Download the current workbook, edit it, and upload it."""
    content = edit_workbook(admin.download_workbook(ADMIN, NOW), edit)
    return admin.upload_workbook("synthetic edited.xlsx", content, ADMIN, NOW)


def add_synthetic_value(bundle: ConfigBundle) -> WorkbookEdit:
    """Return an edit that adds one synthetic value to the roomy sheet."""
    return lambda workbook: append_values(roomy_sheet(workbook, bundle), SYNTHETIC_VALUE)


def apply_new_version(admin: LookupListAdmin, edit: WorkbookEdit) -> str:
    """Upload an edited workbook and apply it; return the new version ID."""
    result = upload_edited(admin, edit)
    assert result.version is not None, result.report.problems
    applied = admin.apply_version(result.version.upload_id, ADMIN, NOW)
    assert applied.version_id is not None
    return applied.version_id


def file_hash(path: Path) -> str:
    """Return a file's SHA-256."""
    return hashlib.sha256(path.read_bytes()).hexdigest()


# ----- Download -----


def test_download_current_has_hidden_metadata(admin: LookupListAdmin, bundle: ConfigBundle) -> None:
    """The downloaded seed has a hidden _metadata sheet with the version and who exported it."""
    workbook = load_workbook(BytesIO(admin.download_workbook(ADMIN, NOW)))
    metadata = workbook[METADATA_SHEET]
    entries = {row[0]: row[1] for row in metadata.iter_rows(max_col=2, values_only=True)}
    assert metadata.sheet_state == "hidden"
    assert entries["version_id"] == bundle.lookup_lists.version
    assert entries["exported_by"] == ADMIN
    assert entries["exported_at"] == NOW.isoformat()


# ----- Upload checks -----


def test_unchanged_upload_passes_with_length_warnings(
    admin: LookupListAdmin, bundle: ConfigBundle
) -> None:
    """Seed values longer than their field's limit are warnings; the upload still passes."""
    result = upload_edited(admin, lambda workbook: None)
    expected_warnings = sum(
        1
        for rule in bundle.field_rules.fields
        if rule.lookup_sheet is not None and rule.max_length is not None
        for value in bundle.lookup_lists.sheets[rule.lookup_sheet]
        if len(value) > rule.max_length
    )
    assert expected_warnings > 0
    assert result.report.problems == []
    assert len(result.report.warnings) == expected_warnings
    assert {w.code for w in result.report.warnings} == {lookup_checks.TOO_LONG}
    assert result.version is not None and result.version.status == "pending"
    assert result.report.changes == []


def test_over_length_value_is_warning_not_rejection(
    admin: LookupListAdmin, bundle: ConfigBundle
) -> None:
    """A new value over the field's limit is a warning, and the version can be applied."""
    rule = roomy_rule(bundle)
    assert rule.max_length is not None
    long_value = "S" * (rule.max_length + 1)
    result = upload_edited(
        admin, lambda workbook: append_values(roomy_sheet(workbook, bundle), long_value)
    )
    warnings = [w for w in result.report.warnings if w.sheet == rule.lookup_sheet]
    assert [w.code for w in warnings] == [lookup_checks.TOO_LONG]
    assert result.version is not None
    admin.apply_version(result.version.upload_id, ADMIN, NOW)
    assert rule.lookup_sheet is not None
    assert long_value in admin.find_current_lookup_lists().sheets[rule.lookup_sheet]


def set_metadata_version(workbook: Workbook, version_id: str) -> None:
    """Overwrite the version_id in the _metadata sheet."""
    sheet = workbook[METADATA_SHEET]
    for row_number, row in enumerate(sheet.iter_rows(max_col=1, values_only=True), start=1):
        if row[0] == "version_id":
            sheet.cell(row=row_number, column=2, value=version_id)


def first_used_sheet(workbook: Workbook, bundle: ConfigBundle) -> Worksheet:
    """Return the sheet of the first field that has a lookup sheet."""
    rule = next(rule for rule in bundle.field_rules.fields if rule.lookup_sheet is not None)
    assert rule.lookup_sheet is not None
    sheet = workbook[rule.lookup_sheet]
    assert isinstance(sheet, Worksheet)
    return sheet


def rename_used_sheet(workbook: Workbook, bundle: ConfigBundle) -> None:
    """Rename a sheet a field uses."""
    first_used_sheet(workbook, bundle).title = "Synthetic renamed"


def delete_used_sheet(workbook: Workbook, bundle: ConfigBundle) -> None:
    """Delete a sheet a field uses."""
    workbook.remove(first_used_sheet(workbook, bundle))


def replace_header(workbook: Workbook, bundle: ConfigBundle) -> None:
    """Overwrite the header cell of a sheet."""
    roomy_sheet(workbook, bundle).cell(row=1, column=1, value="Synthetic header")


def insert_blank_row(workbook: Workbook, bundle: ConfigBundle) -> None:
    """Insert an empty row between two values."""
    roomy_sheet(workbook, bundle).insert_rows(3)


# Each broken edit and the problem code it must give
BROKEN_UPLOADS: list[tuple[str, Callable[[Workbook, ConfigBundle], None], str]] = [
    (
        "stale version",
        lambda wb, b: set_metadata_version(wb, "v0-synthetic-old"),
        lookup_checks.STALE_VERSION,
    ),
    (
        "metadata missing",
        lambda wb, b: wb.remove(wb[METADATA_SHEET]),
        lookup_checks.METADATA_MISSING,
    ),
    ("renamed used sheet", rename_used_sheet, lookup_checks.USED_SHEET_MISSING),
    ("deleted used sheet", delete_used_sheet, lookup_checks.USED_SHEET_MISSING),
    ("wrong header", replace_header, lookup_checks.HEADER_WRONG),
    ("blank row", insert_blank_row, lookup_checks.BLANK_ROW),
    (
        "leading space",
        lambda wb, b: append_values(roomy_sheet(wb, b), f" {SYNTHETIC_VALUE}"),
        lookup_checks.EXTRA_SPACES,
    ),
    (
        "control character",
        lambda wb, b: append_values(roomy_sheet(wb, b), "Synthetic\tvalue"),
        lookup_checks.CONTROL_CHARACTERS,
    ),
    (
        "case-only duplicate",
        lambda wb, b: append_values(roomy_sheet(wb, b), SYNTHETIC_VALUE, SYNTHETIC_VALUE.upper()),
        lookup_checks.CASE_ONLY_DUPLICATE,
    ),
    (
        "exact duplicate",
        lambda wb, b: append_values(roomy_sheet(wb, b), SYNTHETIC_VALUE, SYNTHETIC_VALUE),
        lookup_checks.DUPLICATE_VALUE,
    ),
    ("number cell", lambda wb, b: append_values(roomy_sheet(wb, b), 42), lookup_checks.NOT_TEXT),
]


@pytest.mark.parametrize(
    ("edit", "code"),
    [(edit, code) for _, edit, code in BROKEN_UPLOADS],
    ids=[name for name, _, _ in BROKEN_UPLOADS],
)
def test_broken_upload_is_rejected_with_its_problem(
    admin: LookupListAdmin,
    bundle: ConfigBundle,
    edit: Callable[[Workbook, ConfigBundle], None],
    code: str,
) -> None:
    """Each broken upload reports its problem, creates no version, and is logged."""
    result = upload_edited(admin, lambda workbook: edit(workbook, bundle))
    assert code in {problem.code for problem in result.report.problems}
    assert result.version is None
    assert admin.list_versions() == []
    event = admin.database.read_admin_events()[-1]
    assert event.details["problems"] == len(result.report.problems)


def test_problem_names_sheet_row_and_field(admin: LookupListAdmin, bundle: ConfigBundle) -> None:
    """A renamed used sheet names the field's column; a value problem names sheet and row."""
    renamed = upload_edited(admin, lambda workbook: rename_used_sheet(workbook, bundle))
    rule = next(rule for rule in bundle.field_rules.fields if rule.lookup_sheet is not None)
    assert rule.column in renamed.report.problems[0].message
    spaced = upload_edited(
        admin, lambda wb: append_values(roomy_sheet(wb, bundle), f"{SYNTHETIC_VALUE} ")
    )
    problem = spaced.report.problems[0]
    sheet_name = roomy_rule(bundle).lookup_sheet
    assert sheet_name is not None
    assert problem.sheet == sheet_name
    assert problem.row == len(bundle.lookup_lists.sheets[sheet_name]) + 2


def test_unreadable_file_is_rejected(admin: LookupListAdmin) -> None:
    """Bytes that are not a workbook give one FILE_UNREADABLE problem."""
    result = admin.upload_workbook("synthetic.xlsx", b"synthetic, not a workbook", ADMIN, NOW)
    assert [problem.code for problem in result.report.problems] == [lookup_checks.FILE_UNREADABLE]
    assert result.report.passed_checks == []
    assert result.version is None


def test_failed_checks_are_not_listed_as_passed(
    admin: LookupListAdmin, bundle: ConfigBundle
) -> None:
    """A stale upload does not show the "based on the current version" check as passed."""
    result = upload_edited(admin, lambda wb: set_metadata_version(wb, "v0-synthetic-old"))
    assert not any("current version" in check for check in result.report.passed_checks)
    assert any("headers present" in check for check in result.report.passed_checks)


# ----- Diff -----


def test_upload_reports_added_and_removed_values(
    admin: LookupListAdmin, bundle: ConfigBundle
) -> None:
    """Adding a value and deleting the last one of a sheet show up in the changes."""
    sheet_name = roomy_rule(bundle).lookup_sheet
    assert sheet_name is not None
    seed_values = bundle.lookup_lists.sheets[sheet_name]

    def edit(workbook: Workbook) -> None:
        sheet = roomy_sheet(workbook, bundle)
        sheet.cell(row=len(seed_values) + 1, column=1, value=SYNTHETIC_VALUE)

    result = upload_edited(admin, edit)
    assert [change.sheet for change in result.report.changes] == [sheet_name]
    change = result.report.changes[0]
    assert (change.added, change.removed) == ([SYNTHETIC_VALUE], [seed_values[-1]])
    assert result.version is not None
    assert (result.version.summary["added_values"], result.version.summary["removed_values"]) == (
        1,
        1,
    )


def test_diff_marks_new_and_removed_sheets() -> None:
    """A sheet only in the upload is new; a sheet only in the current version is removed."""
    current = {"Synthetic list A": ["a1", "a2"], "Synthetic list B": ["b1"]}
    uploaded = {"Synthetic list A": ["a1", "a2"], "Synthetic list C": ["c1"]}
    changes = diff_sheets(current, uploaded)
    assert [(c.sheet, c.is_new_sheet, c.is_removed_sheet) for c in changes] == [
        ("Synthetic list C", True, False),
        ("Synthetic list B", False, True),
    ]
    assert (changes[0].added, changes[1].removed) == (["c1"], ["b1"])


def test_new_unused_sheet_is_stored(admin: LookupListAdmin) -> None:
    """An admin can add a sheet no field uses; it is stored with the version."""

    def add_sheet(workbook: Workbook) -> None:
        sheet = workbook.create_sheet("Synthetic new list")
        sheet.append([HEADER])
        sheet.append([SYNTHETIC_VALUE])

    version_id = apply_new_version(admin, add_sheet)
    assert admin.load_lookup_lists(version_id).sheets["Synthetic new list"] == [SYNTHETIC_VALUE]


# ----- Apply, restore, seed -----


def test_apply_gives_v1_then_v2(admin: LookupListAdmin, bundle: ConfigBundle) -> None:
    """The first apply after the seed is v1, the next is v2; the old one is superseded."""
    assert apply_new_version(admin, add_synthetic_value(bundle)) == "v1"
    second = upload_edited(
        admin, lambda wb: append_values(roomy_sheet(wb, bundle), "Synthetic second")
    )
    assert second.report.based_on_version == "v1"
    assert second.version is not None
    assert admin.apply_version(second.version.upload_id, ADMIN, NOW).version_id == "v2"
    statuses = [(version.version_id, version.status) for version in admin.list_versions()]
    assert statuses == [("v1", "superseded"), ("v2", "current")]
    assert admin.find_current_lookup_lists().version == "v2"


def test_restore_makes_an_older_version_current(
    admin: LookupListAdmin, bundle: ConfigBundle
) -> None:
    """Restoring v1 after v2 makes v1 current again under its own name, and is logged."""
    apply_new_version(admin, add_synthetic_value(bundle))
    apply_new_version(admin, lambda wb: append_values(roomy_sheet(wb, bundle), "Synthetic two"))
    restored = admin.apply_version("v1", ADMIN, NOW)
    assert (restored.version_id, restored.status) == ("v1", "current")
    assert admin.find_current_lookup_lists().version == "v1"
    events = [event.event for event in admin.database.read_admin_events()]
    assert events[-1] == "lookup_version_restored"
    assert events.count("lookup_version_applied") == 2


def test_pending_upload_is_stale_after_another_apply(
    admin: LookupListAdmin, bundle: ConfigBundle
) -> None:
    """Two uploads from the same download: once one is applied, the other is refused."""
    downloaded = admin.download_workbook(ADMIN, NOW)
    first = admin.upload_workbook(
        "one.xlsx", edit_workbook(downloaded, add_synthetic_value(bundle)), ADMIN, NOW
    )
    second = admin.upload_workbook("two.xlsx", downloaded, ADMIN, NOW)
    assert first.version is not None and second.version is not None
    admin.apply_version(first.version.upload_id, ADMIN, NOW)
    with pytest.raises(AdminError) as error:
        admin.apply_version(second.version.upload_id, ADMIN, NOW)
    assert error.value.code == "UPLOAD_STALE"


def test_apply_refuses_current_and_unknown_versions(
    admin: LookupListAdmin, bundle: ConfigBundle
) -> None:
    """Applying the current version or an unknown one gives a plain error."""
    apply_new_version(admin, add_synthetic_value(bundle))
    with pytest.raises(AdminError) as current:
        admin.apply_version("v1", ADMIN, NOW)
    with pytest.raises(AdminError) as unknown:
        admin.apply_version("v99", ADMIN, NOW)
    assert (current.value.code, unknown.value.code) == (
        "VERSION_ALREADY_CURRENT",
        "VERSION_NOT_FOUND",
    )


def test_seed_is_never_written(
    admin: LookupListAdmin, bundle: ConfigBundle, tmp_path: Path
) -> None:
    """Downloads, uploads, applies and restores leave the seed file unchanged."""
    seed_path = bundle.config_dir / "lookup_lists.xlsx"
    before = file_hash(seed_path)
    apply_new_version(admin, add_synthetic_value(bundle))
    apply_new_version(admin, lambda wb: append_values(roomy_sheet(wb, bundle), "Synthetic two"))
    admin.apply_version("v1", ADMIN, NOW)
    admin.download_workbook(ADMIN, NOW, bundle.lookup_lists.version)
    assert file_hash(seed_path) == before
    stored = sorted(path.name for path in (tmp_path / "storage" / "lookup_lists").iterdir())
    assert stored == ["upl_0001.xlsx", "upl_0002.xlsx"]


def test_new_runs_pick_up_the_current_version(admin: LookupListAdmin, bundle: ConfigBundle) -> None:
    """Before any apply new runs get the seed; afterwards the current version. Old ones stay."""
    assert admin.find_current_lookup_lists() == bundle.lookup_lists
    apply_new_version(admin, add_synthetic_value(bundle))
    current = admin.find_current_lookup_lists()
    sheet_name = roomy_rule(bundle).lookup_sheet
    assert sheet_name is not None
    assert (current.version, current.metadata.based_on_version) == (
        "v1",
        bundle.lookup_lists.version,
    )
    assert SYNTHETIC_VALUE in current.sheets[sheet_name]
    # A run that started on the seed keeps the seed
    assert admin.load_lookup_lists(bundle.lookup_lists.version) == bundle.lookup_lists


def test_download_of_a_version_names_that_version(
    admin: LookupListAdmin, bundle: ConfigBundle
) -> None:
    """After v1 is applied, the downloaded workbook says v1, based on the seed."""
    apply_new_version(admin, add_synthetic_value(bundle))
    workbook = load_workbook(BytesIO(admin.download_workbook(ADMIN, NOW)))
    entries = {
        row[0]: row[1] for row in workbook[METADATA_SHEET].iter_rows(max_col=2, values_only=True)
    }
    assert (entries["version_id"], entries["based_on_version"]) == (
        "v1",
        bundle.lookup_lists.version,
    )
