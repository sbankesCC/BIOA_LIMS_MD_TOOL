"""Broken-config tests: each copy of the config has one problem, and loading names it.

Every case edits a temp copy of the client config in code, so no client values are written here.
"""

from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from app.config.bundle import ConfigBundle, load_config_bundle
from app.config.errors import ConfigError
from tests.conftest import edit_yaml


def assert_config_error(config_dir: Path, expected_text: str) -> None:
    """Load the config and check it fails with a problem containing `expected_text`."""
    with pytest.raises(ConfigError) as caught:
        load_config_bundle(config_dir)
    assert any(expected_text in problem for problem in caught.value.problems), caught.value


def first_lookup_field(fields: list[dict[str, Any]]) -> dict[str, Any]:
    """Return the first field rule that has a lookup sheet."""
    return next(field for field in fields if field.get("lookup_sheet"))


# ----- Missing files and folders -----


def test_missing_config_folder_is_reported(tmp_path: Path) -> None:
    """A config folder that does not exist names CLIENT_CONFIG_DIR."""
    assert_config_error(tmp_path / "nowhere", "CLIENT_CONFIG_DIR")


def test_missing_prompt_file_is_reported(config_copy: Path) -> None:
    """A missing prompt file is reported as not found."""
    (config_copy / "prompts" / "duplicate_judge.md").unlink()
    with pytest.raises(ConfigError, match="duplicate_judge.md: file not found"):
        load_config_bundle(config_copy)


def test_invalid_yaml_is_reported(config_copy: Path) -> None:
    """A YAML syntax error names the file."""
    (config_copy / "issue_catalogue.yaml").write_text("version: [unclosed", encoding="utf-8")
    assert_config_error(config_copy, "issue_catalogue.yaml: not valid YAML")


# ----- Field rules -----


def test_duplicate_column_is_reported(config_copy: Path) -> None:
    """Two fields with the same column name are rejected."""

    def change(data: Any) -> None:
        data["fields"][1]["column"] = data["fields"][0]["column"]

    edit_yaml(config_copy / "field_rules.yaml", change)
    assert_config_error(config_copy, "is used twice")


def test_unknown_role_is_reported(config_copy: Path) -> None:
    """A role outside extracted / blank_for_now / system is rejected."""
    edit_yaml(config_copy / "field_rules.yaml", lambda data: data["fields"][0].update(role="guess"))
    assert_config_error(config_copy, "fields.0.role")


def test_extracted_field_without_guidance_is_reported(config_copy: Path) -> None:
    """An extracted field needs extraction guidance for the prompt."""
    edit_yaml(
        config_copy / "field_rules.yaml",
        lambda data: data["fields"][0].pop("extraction_guidance"),
    )
    assert_config_error(config_copy, "has no extraction_guidance")


def test_no_record_name_field_is_reported(config_copy: Path) -> None:
    """A config where no field has is_record_name is rejected."""

    def change(data: Any) -> None:
        for field in data["fields"]:
            field.pop("is_record_name", None)

    edit_yaml(config_copy / "field_rules.yaml", change)
    assert_config_error(config_copy, "exactly one field must have is_record_name: true, but 0 do")


def test_two_record_name_fields_are_reported(config_copy: Path) -> None:
    """A config where two fields have is_record_name is rejected."""

    def change(data: Any) -> None:
        unflagged = next(field for field in data["fields"] if not field.get("is_record_name"))
        unflagged.update(is_record_name=True, sent_to_duplicate_judge=True)

    edit_yaml(config_copy / "field_rules.yaml", change)
    assert_config_error(config_copy, "exactly one field must have is_record_name: true, but 2 do")


def test_record_name_not_sent_to_judge_is_reported(config_copy: Path) -> None:
    """The record name field must also be sent to the duplicate judge."""

    def change(data: Any) -> None:
        name_field = next(field for field in data["fields"] if field.get("is_record_name"))
        name_field.pop("sent_to_duplicate_judge", None)

    edit_yaml(config_copy / "field_rules.yaml", change)
    assert_config_error(config_copy, "has is_record_name but not sent_to_duplicate_judge")


def test_conditional_field_outside_any_group_is_reported(config_copy: Path) -> None:
    """A conditional field must belong to a condition group."""
    edit_yaml(config_copy / "field_rules.yaml", lambda data: data.update(condition_groups={}))
    assert_config_error(config_copy, "is conditional but in no condition group")


def test_condition_group_with_unknown_field_is_reported(config_copy: Path) -> None:
    """A condition group naming a field that does not exist is rejected."""

    def change(data: Any) -> None:
        group = next(g for g in data["condition_groups"].values() if g["rule"] == "both_or_neither")
        group["fields"] = [*group["fields"][:-1], "no_such_field"]

    edit_yaml(config_copy / "field_rules.yaml", change)
    assert_config_error(config_copy, "unknown field 'no_such_field'")


def test_field_order_must_match_template(config_copy: Path) -> None:
    """Swapping two fields makes the field rules disagree with the DFD template headers."""

    def change(data: Any) -> None:
        data["fields"][0], data["fields"][1] = data["fields"][1], data["fields"][0]

    edit_yaml(config_copy / "field_rules.yaml", change)
    assert_config_error(config_copy, "the headers do not match the columns")


def test_default_not_on_lookup_sheet_is_reported(config_copy: Path) -> None:
    """A lookup field's default must be one of the sheet's values."""
    edit_yaml(
        config_copy / "field_rules.yaml",
        lambda data: first_lookup_field(data["fields"]).update(default="synthetic-not-a-value"),
    )
    assert_config_error(config_copy, "is not on the")


# ----- DFD template -----


def test_template_with_extra_rows_is_reported(config_copy: Path) -> None:
    """The template may hold only its header row."""
    path = config_copy / "dfd_template.xlsx"
    workbook = load_workbook(path)
    workbook["Sheet1"]["A2"] = "synthetic record"
    workbook.save(path)
    assert_config_error(config_copy, "must have only the header row")


# ----- Lookup lists -----


def test_missing_lookup_sheet_is_reported(config_copy: Path, bundle: ConfigBundle) -> None:
    """A field whose lookup sheet is missing from the workbook is rejected."""
    sheet_name = next(f.lookup_sheet for f in bundle.field_rules.fields if f.lookup_sheet)
    path = config_copy / "lookup_lists.xlsx"
    workbook = load_workbook(path)
    del workbook[sheet_name]
    workbook.save(path)
    assert_config_error(config_copy, "is missing")


def test_wrong_lookup_header_is_reported(config_copy: Path, bundle: ConfigBundle) -> None:
    """Each lookup sheet must have 'Lookup Value' in A1."""
    sheet_name = next(iter(bundle.lookup_lists.sheets))
    path = config_copy / "lookup_lists.xlsx"
    workbook = load_workbook(path)
    workbook[sheet_name]["A1"] = "Value"
    workbook.save(path)
    assert_config_error(config_copy, "cell A1 must be 'Lookup Value'")


def test_blank_row_inside_lookup_list_is_reported(config_copy: Path, bundle: ConfigBundle) -> None:
    """A gap between lookup values is rejected."""
    sheet_name = next(name for name, values in bundle.lookup_lists.sheets.items() if values)
    path = config_copy / "lookup_lists.xlsx"
    workbook = load_workbook(path)
    worksheet = workbook[sheet_name]
    worksheet.insert_rows(2)
    workbook.save(path)
    assert_config_error(config_copy, "row 2 must be a text value")


def test_missing_metadata_version_is_reported(config_copy: Path) -> None:
    """The `_metadata` sheet must give a version_id."""
    path = config_copy / "lookup_lists.xlsx"
    workbook = load_workbook(path)
    workbook["_metadata"]["B1"] = None
    workbook.save(path)
    assert_config_error(config_copy, "has no version_id")


# ----- Issue catalogue -----


def test_unknown_severity_is_reported(config_copy: Path) -> None:
    """Severities are error, warning or info only."""

    def change(data: Any) -> None:
        first_code = next(iter(data["record_issues"]))
        data["record_issues"][first_code]["severity"] = "critical"

    edit_yaml(config_copy / "issue_catalogue.yaml", change)
    assert_config_error(config_copy, "severity")


def test_individual_review_on_error_is_reported(config_copy: Path) -> None:
    """Only warnings can be accepted, so individual_review on an error is rejected."""

    def change(data: Any) -> None:
        code = next(c for c, d in data["record_issues"].items() if d["severity"] == "error")
        data["record_issues"][code]["individual_review"] = True

    edit_yaml(config_copy / "issue_catalogue.yaml", change)
    assert_config_error(config_copy, "has individual_review but is not a warning")


def test_unknown_message_placeholder_is_reported(config_copy: Path) -> None:
    """A message placeholder the code cannot fill is rejected."""

    def change(data: Any) -> None:
        first_code = next(iter(data["record_issues"]))
        data["record_issues"][first_code]["message"] = "Check {mystery}."

    edit_yaml(config_copy / "issue_catalogue.yaml", change)
    assert_config_error(config_copy, "uses {mystery}")


def test_condition_group_code_missing_from_catalogue_is_reported(
    config_copy: Path, bundle: ConfigBundle
) -> None:
    """A condition group's issue code must exist in the catalogue."""
    group_code = next(iter(bundle.field_rules.condition_groups.values())).issue_code
    edit_yaml(
        config_copy / "issue_catalogue.yaml",
        lambda data: data["record_issues"].pop(group_code),
    )
    assert_config_error(config_copy, f"uses issue code {group_code}")


# ----- Prompts -----


def test_wrong_prompt_header_is_reported(config_copy: Path) -> None:
    """A prompt whose first line is not the header comment is rejected."""
    path = config_copy / "prompts" / "extraction.md"
    text = path.read_text(encoding="utf-8")
    path.write_text("# No header\n" + text, encoding="utf-8")
    assert_config_error(config_copy, "prompts/extraction.md: the first line must be")


def test_prompt_header_name_must_match_file(config_copy: Path) -> None:
    """The name in the header must be the file name."""
    path = config_copy / "prompts" / "duplicate_judge.md"
    text = path.read_text(encoding="utf-8")
    path.write_text(text.replace("prompt: duplicate_judge", "prompt: judge", 1), encoding="utf-8")
    assert_config_error(config_copy, "the header names 'judge'")


def test_missing_prompt_placeholder_is_reported(config_copy: Path) -> None:
    """A placeholder the code fills must appear in the prompt body."""
    path = config_copy / "prompts" / "duplicate_judge.md"
    text = path.read_text(encoding="utf-8")
    path.write_text(text.replace("{candidates}", "(none)"), encoding="utf-8")
    assert_config_error(config_copy, "placeholder {candidates} is missing")


def test_unknown_prompt_placeholder_is_reported(config_copy: Path) -> None:
    """A placeholder the code does not fill is rejected."""
    path = config_copy / "prompts" / "duplicate_judge.md"
    text = path.read_text(encoding="utf-8")
    path.write_text(text + "\n{mystery}\n", encoding="utf-8")
    assert_config_error(config_copy, "placeholder {mystery} is unknown")


def test_placeholders_inside_html_comments_are_ignored(config_copy: Path) -> None:
    """Braces in an HTML comment are notes for developers, not placeholders."""
    path = config_copy / "prompts" / "duplicate_judge.md"
    text = path.read_text(encoding="utf-8")
    path.write_text(text + "\n<!-- {mystery} is only a note -->\n", encoding="utf-8")
    assert load_config_bundle(config_copy).prompts["duplicate_judge"]


# ----- Reference data mapping -----


def test_reference_mapping_without_name_column_is_reported(config_copy: Path) -> None:
    """The name column is required for the duplicate check."""
    edit_yaml(
        config_copy / "reference_data_mapping.yaml",
        lambda data: data["columns"].pop("name"),
    )
    assert_config_error(config_copy, "columns.name")
