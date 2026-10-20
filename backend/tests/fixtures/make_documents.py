"""Synthetic test documents (Word, text PDF, scanned PDF, Excel, loader, broken file).

Every file is labeled `synthetic` in its name and content; none holds real client data.
Tests generate them into temp folders; run `python make_documents.py <folder>` to look at them.
"""

import sys
from pathlib import Path

import docx
from docx.document import Document as WordDocument
from fpdf import FPDF
from openpyxl import Workbook
from PIL import Image, ImageDraw

# File names of the generated documents
WORD_FILE = "synthetic_method.docx"
LARGE_TABLE_WORD_FILE = "synthetic_large_table.docx"
TEXT_PDF_FILE = "synthetic_study_plan.pdf"
SCANNED_PDF_FILE = "synthetic_scanned_worksheet.pdf"
WORKBOOK_FILE = "synthetic_reagent_list.xlsx"
LOADER_FILE = "synthetic_loader.xlsx"
BROKEN_FILE = "synthetic_broken.docx"

# The synthetic reagent table used in the Word method (header row first)
REAGENT_TABLE = [
    ["Item", "Vendor", "Cat. number", "Keep at"],
    ["Synthetic Buffer A", "Synthetic Vendor One", "SYN-1001", "synthetic cold room"],
    ["Synthetic Buffer B", "Synthetic Vendor One", "SYN-1002", "synthetic cold room"],
    ["Synthetic Diluent C", "Synthetic Vendor Two", "SYN-2001", "synthetic shelf"],
    ["Synthetic Plate D", "Synthetic Vendor Two", "SYN-2002", "synthetic shelf"],
]

# Rows in the large Word table and the large sheet (enough to be split)
LARGE_ROW_COUNT = 120


def make_word_document(path: Path) -> None:
    """Write a Word method with headings, paragraphs and one reagent table."""
    document = docx.Document()
    document.add_heading("Synthetic Method SYN-001", level=0)
    document.add_heading("1 Purpose", level=1)
    document.add_paragraph("This synthetic method describes a made-up test procedure.")
    document.add_heading("6 Materials", level=1)
    document.add_heading("6.2 Reagents", level=2)
    document.add_paragraph("The following synthetic reagents are used.")
    add_word_table(document, REAGENT_TABLE)
    document.add_paragraph("Synthetic reagents are prepared on the day of use.")
    document.add_heading("6.3 Handling", level=2)
    document.add_paragraph("Synthetic handling notes for the made-up procedure.")
    document.save(str(path))


def make_large_table_word_document(path: Path) -> None:
    """Write a Word file with one very large table, which must be split."""
    document = docx.Document()
    document.add_heading("Synthetic large table", level=1)
    rows = [REAGENT_TABLE[0]] + [
        [f"Synthetic Item {number}", "Synthetic Vendor", f"SYN-{number:04d}", "synthetic shelf"]
        for number in range(1, LARGE_ROW_COUNT + 1)
    ]
    add_word_table(document, rows)
    document.save(str(path))


def add_word_table(document: WordDocument, rows: list[list[str]]) -> None:
    """Add a table with the given rows to a Word document."""
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for row_index, row in enumerate(rows):
        for column_index, text in enumerate(row):
            table.cell(row_index, column_index).text = text


def make_text_pdf(path: Path) -> None:
    """Write a two-page PDF with a text layer on each page."""
    pdf = FPDF()
    pdf.set_font("Helvetica", size=12)
    for page_text in (
        "Synthetic study plan, page one. Synthetic Buffer A is used.",
        "Synthetic study plan, page two. Synthetic Diluent C is used.",
    ):
        pdf.add_page()
        pdf.multi_cell(0, 10, page_text)
    pdf.output(str(path))


def make_scanned_pdf(path: Path) -> None:
    """Write a one-page PDF that is only an image (no text layer), like a scanned worksheet."""
    image = Image.new("RGB", (850, 1100), "white")
    drawing = ImageDraw.Draw(image)
    drawing.text((60, 60), "Synthetic scanned worksheet", fill="black")
    drawing.text((60, 100), "Synthetic Buffer B, lot SYN-LOT-01", fill="black")
    image.save(path, format="PDF")


def make_workbook(path: Path) -> None:
    """Write a workbook with a small reagent sheet and a large sheet that must be split."""
    workbook = Workbook()
    small_sheet = workbook.active
    assert small_sheet is not None
    small_sheet.title = "Reagents"
    for row in REAGENT_TABLE:
        small_sheet.append(row)
    large_sheet = workbook.create_sheet("Large Sheet")
    large_sheet.append(REAGENT_TABLE[0])
    for number in range(1, LARGE_ROW_COUNT + 1):
        large_sheet.append([f"Synthetic Item {number}", "Synthetic Vendor", number])
    workbook.save(path)


def make_loader_workbook(path: Path, dfd_headers: list[str]) -> None:
    """Write a workbook shaped like a DFD loader; the headers come from the client config."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(dfd_headers)
    sheet.append(["synthetic"] * len(dfd_headers))
    workbook.save(path)


def make_broken_file(path: Path) -> None:
    """Write a file with a Word extension that is not a Word file."""
    path.write_bytes(b"synthetic: this is not a real Word document")


def make_all_documents(folder: Path, dfd_headers: list[str]) -> dict[str, Path]:
    """Write every synthetic document into `folder` and return their paths by file name."""
    folder.mkdir(parents=True, exist_ok=True)
    paths = {
        name: folder / name
        for name in (
            WORD_FILE,
            LARGE_TABLE_WORD_FILE,
            TEXT_PDF_FILE,
            SCANNED_PDF_FILE,
            WORKBOOK_FILE,
            LOADER_FILE,
            BROKEN_FILE,
        )
    }
    make_word_document(paths[WORD_FILE])
    make_large_table_word_document(paths[LARGE_TABLE_WORD_FILE])
    make_text_pdf(paths[TEXT_PDF_FILE])
    make_scanned_pdf(paths[SCANNED_PDF_FILE])
    make_workbook(paths[WORKBOOK_FILE])
    make_loader_workbook(paths[LOADER_FILE], dfd_headers)
    make_broken_file(paths[BROKEN_FILE])
    return paths


if __name__ == "__main__":
    # Manual use: the loader file gets placeholder headers, since no client config is loaded here
    make_all_documents(Path(sys.argv[1]), ["synthetic column one", "synthetic column two"])
